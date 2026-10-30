"""prepare-excel command and the end-to-end run through shaclplay and sempyro."""

import shutil

from openpyxl import load_workbook
from prep_helpers import CLASS_ANSWERS, fill_reported_gaps

from metadata_automation.cli import derive_fdp_excel, prepare_excel, sempyro, shaclplay
from metadata_automation.excel_prep.fdp import derive_fdp_workbook
from metadata_automation.excel_prep.workbook import class_sheet_names, data_rows, headers, text


def test_plain_metadata_gets_all_classes(plain_metadata, tmp_path):
    out = tmp_path / "prepared.xlsx"
    derive_fdp_workbook(plain_metadata, out)
    wb = load_workbook(out)
    ws = wb["classes"]
    cols = headers(ws)
    names = [text(ws, r, cols["sheet_name"]) for r in data_rows(ws, cols, ("sheet_name",))]
    assert names == class_sheet_names(wb)
    assert "Resource" in names and "CatalogueRecord" in names


def test_cli_prints_report(runner, plain_small, tmp_path):
    out = tmp_path / "prepared.xlsx"
    result = runner.invoke(derive_fdp_excel, ["-i", str(plain_small), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert "Required cells to fill in" in result.output
    assert out.exists()


def test_cli_in_place_flag(runner, plain_small, tmp_path):
    path = tmp_path / "plain.xlsx"
    shutil.copy(plain_small, path)
    result = runner.invoke(prepare_excel, ["-i", str(path), "--in-place"])
    assert result.exit_code == 0, result.output
    assert "Prefix check" in result.output
    assert f"Prepared Excel written to {path}" in result.output
    assert "prefixes" in load_workbook(path).sheetnames


def test_cli_in_place_rejects_output(runner, plain_small, tmp_path):
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(tmp_path / "out.xlsx"), "--in-place"])
    assert result.exit_code == 1
    assert "either an output path (-o) or --in-place" in result.output


def test_cli_needs_output_or_in_place(runner, plain_small):
    result = runner.invoke(prepare_excel, ["-i", str(plain_small)])
    assert result.exit_code == 1
    assert "either an output path (-o) or --in-place" in result.output


def test_cli_prepare_to_new_file_leaves_source_alone(runner, plain_small, tmp_path):
    out = tmp_path / "prepared.xlsx"
    before = plain_small.read_bytes()
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert plain_small.read_bytes() == before
    assert "Resource" not in load_workbook(out).sheetnames


def test_cli_derive_fdp_excel_requires_output(runner, plain_small):
    result = runner.invoke(derive_fdp_excel, ["-i", str(plain_small)])
    assert result.exit_code == 2
    assert "--output-excel" in result.output


def test_cli_rejects_same_output(runner, plain_small):
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(plain_small)])
    assert result.exit_code == 1
    assert "must differ" in result.output


def test_end_to_end_shaclplay_and_sempyro(runner, second_pass, tmp_path):
    _, _, path = second_pass
    fill_reported_gaps(path)
    final = tmp_path / "final.xlsx"
    report = derive_fdp_workbook(path, final)
    assert report.gaps == []

    result = runner.invoke(shaclplay, ["-i", str(final), "-o", str(tmp_path / "shaclplay")])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "shaclplay" / "SHACL-catalog.xlsx").exists()

    imports = tmp_path / "imports.yaml"
    header = (
        "    from pydantic import AnyHttpUrl, ConfigDict, Field\n"
        "    from rdflib.namespace import DCAT, DCTERMS, FOAF\n"
        "    from sempyro import LiteralField, RDFModel\n"
    )
    imports.write_text("".join(f"hri-{name}: |\n{header}\n" for name in CLASS_ANSWERS))
    result = runner.invoke(
        sempyro,
        [
            "-i",
            str(final),
            "--linkml-output-path",
            str(tmp_path / "linkml"),
            "--sempyro-output-path",
            str(tmp_path / "sempyro"),
            "--imports-path",
            str(imports),
        ],
    )
    assert result.exit_code == 0, result.output
    assert "Successfully generated: 5 classes" in result.output
