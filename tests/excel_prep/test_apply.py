"""Applying a plan: equal to the preparation, safe against a changed source, and logged."""

import shutil

import pytest
from openpyxl import load_workbook
from prep_helpers import CLASS_ANSWERS, answer_classes, cell, find_row, sha

from metadata_automation.cli import apply_plan, plan_excel
from metadata_automation.excel_prep.apply import PlanError
from metadata_automation.excel_prep.apply import apply_plan as apply_plan_to_workbook
from metadata_automation.excel_prep.plan import compute_plan, default_plan_path, write_plan
from metadata_automation.excel_prep.preparer import prepare_workbook
from metadata_automation.excel_prep.workbook import text


def values(path):
    """All cell values per sheet, as stripped text, to compare workbooks."""
    wb = load_workbook(path, rich_text=True)
    return {
        ws.title: [[text(ws, r, c) for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
        for ws in wb.worksheets
    }


@pytest.fixture
def source(plain_small, tmp_path):
    path = tmp_path / "model.xlsx"
    shutil.copy(plain_small, path)
    return path


@pytest.fixture
def answered(source):
    """Prepared and with the class questions answered, so the plan has cell edits too."""
    prepare_workbook(source)
    answer_classes(source, {k: v for k, v in CLASS_ANSWERS.items() if k != "Resource"})
    return source


def make_plan(path, namespace=None):
    plan, _ = compute_plan(path, namespace)
    plan_path = default_plan_path(path)
    write_plan(plan, plan_path)
    return plan, plan_path


class TestApply:
    @pytest.mark.parametrize("fixture", ["source", "answered"])
    def test_applying_the_plan_equals_preparing(self, request, fixture, tmp_path):
        path = request.getfixturevalue(fixture)
        expected = tmp_path / "expected.xlsx"
        prepare_workbook(path, expected)
        _, plan_path = make_plan(path)

        apply_plan_to_workbook(path, plan_path)

        assert values(path) == values(expected)

    def test_a_second_plan_has_no_edits(self, answered):
        _, plan_path = make_plan(answered)
        apply_plan_to_workbook(answered, plan_path)
        plan, _ = compute_plan(answered)
        assert plan.edits == []
        assert plan.gaps

    def test_plan_is_deleted_after_updating_in_place(self, source):
        _, plan_path = make_plan(source)
        apply_plan_to_workbook(source, plan_path)
        assert not plan_path.exists()

    def test_keep_plan(self, source):
        _, plan_path = make_plan(source)
        apply_plan_to_workbook(source, plan_path, keep_plan=True)
        assert plan_path.exists()

    def test_output_file_leaves_source_and_plan_alone(self, source, tmp_path):
        _, plan_path = make_plan(source)
        before = sha(source)
        out = tmp_path / "out.xlsx"
        apply_plan_to_workbook(source, plan_path, out)
        assert sha(source) == before
        assert plan_path.exists()
        assert "classes" in load_workbook(out).sheetnames

    def test_edit_marked_no_is_skipped(self, answered):
        plan, plan_path = make_plan(answered)
        wb = load_workbook(plan_path)
        edits = wb["Edits"]
        target = next(r for r in range(2, edits.max_row + 1) if edits.cell(row=r, column=2).value == "set cell")
        edits.cell(row=target, column=1, value="no")
        column, key = edits.cell(row=target, column=6).value, edits.cell(row=target, column=4).value
        sheet = edits.cell(row=target, column=3).value
        wb.save(plan_path)

        result = apply_plan_to_workbook(answered, plan_path)

        assert len(result.skipped) == 1
        ws = load_workbook(answered)[sheet]
        row = find_row(ws, "Property URI", key)
        assert not cell(ws, row, column)

    def test_proposed_value_can_be_changed(self, answered):
        _, plan_path = make_plan(answered)
        wb = load_workbook(plan_path)
        edits = wb["Edits"]
        for r in range(2, edits.max_row + 1):
            if edits.cell(row=r, column=2).value == "set cell" and edits.cell(row=r, column=6).value == "SHACL_sh:node":
                sheet, key = edits.cell(row=r, column=3).value, edits.cell(row=r, column=4).value
                edits.cell(row=r, column=8, value="custom:Shape")
                break
        wb.save(plan_path)

        apply_plan_to_workbook(answered, plan_path)

        ws = load_workbook(answered)[sheet]
        assert cell(ws, find_row(ws, "Property URI", key), "SHACL_sh:node") == "custom:Shape"

    def test_gap_values_are_written_and_the_gap_is_closed(self, answered):
        _, plan_path = make_plan(answered)
        wb = load_workbook(plan_path)
        gaps = wb["Gaps"]
        for r in range(2, gaps.max_row + 1):
            if (
                gaps.cell(row=r, column=1).value == "classes"
                and gaps.cell(row=r, column=4).value == "SeMPyRO_annotations_IRI"
            ):
                key = gaps.cell(row=r, column=2).value
                gaps.cell(row=r, column=5, value="https://example.org/x")
                break
        wb.save(plan_path)

        result = apply_plan_to_workbook(answered, plan_path)

        ws = load_workbook(answered)["classes"]
        assert cell(ws, find_row(ws, "sheet_name", key), "SeMPyRO_annotations_IRI") == "https://example.org/x"
        assert len(result.filled) == 1
        plan, _ = compute_plan(answered)
        assert not any(g.key == key and g.column == "SeMPyRO_annotations_IRI" for g in plan.gaps)

    def test_rows_inserted_after_planning_are_tolerated_when_the_hash_is_ignored(self, answered, tmp_path):
        expected = tmp_path / "expected.xlsx"
        prepare_workbook(answered, expected)
        _, plan_path = make_plan(answered)
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        ws.insert_rows(2)
        ws.cell(row=2, column=1, value="inserted by the data team")
        wb.save(answered)

        with pytest.raises(PlanError, match="--ignore-source-hash"):
            apply_plan_to_workbook(answered, plan_path)
        apply_plan_to_workbook(answered, plan_path, check_hash=False)

        # Every planned edit landed on its own property, not on the shifted row number.
        ws = load_workbook(answered)["Catalog"]
        reference = load_workbook(expected)["Catalog"]
        for uri in ("dct:publisher", "dcat:dataset"):
            for column in ("SHACL_sh:node", "SHACL_dash:viewer", "SHACL_dash:editor", "Range"):
                assert cell(ws, find_row(ws, "Property URI", uri), column) == cell(
                    reference, find_row(reference, "Property URI", uri), column
                )


class TestRefusals:
    def test_changed_source_is_refused_and_nothing_is_written(self, answered):
        _, plan_path = make_plan(answered)
        wb = load_workbook(answered)
        wb["Catalog"]["A2"] = "edited by the data team"
        wb.save(answered)
        before = sha(answered)

        with pytest.raises(PlanError, match="changed since the plan was made"):
            apply_plan_to_workbook(answered, plan_path)
        assert sha(answered) == before
        assert plan_path.exists()

    def test_edit_that_does_not_fit_writes_nothing(self, answered):
        _, plan_path = make_plan(answered)
        wb = load_workbook(plan_path)
        edits = wb["Edits"]
        for r in range(2, edits.max_row + 1):
            if edits.cell(row=r, column=2).value == "set cell":
                edits.cell(row=r, column=7, value="something else")  # the expected current value
                break
        wb.save(plan_path)
        before = sha(answered)

        with pytest.raises(PlanError, match="the plan expected"):
            apply_plan_to_workbook(answered, plan_path)
        assert sha(answered) == before

    def test_skipping_a_sheet_that_later_edits_need_is_reported(self, source):
        _, plan_path = make_plan(source)
        wb = load_workbook(plan_path)
        edits = wb["Edits"]
        for r in range(2, edits.max_row + 1):
            if edits.cell(row=r, column=2).value == "add sheet" and edits.cell(row=r, column=3).value == "classes":
                edits.cell(row=r, column=1, value="no")
        wb.save(plan_path)
        before = sha(source)

        with pytest.raises(PlanError, match="sheet 'classes' does not exist"):
            apply_plan_to_workbook(source, plan_path)
        assert sha(source) == before


class TestCommands:
    def test_plan_then_apply(self, runner, source):
        assert runner.invoke(plan_excel, ["-i", str(source)]).exit_code == 0
        plan_path = default_plan_path(source)
        log = source.parent / "apply.log"

        result = runner.invoke(apply_plan, ["-p", str(plan_path), "--log", str(log)])

        assert result.exit_code == 0, result.output
        assert "Applied (" in result.output
        assert log.read_text().startswith("Applied (")
        assert not plan_path.exists()
        assert "prefixes" in load_workbook(source).sheetnames

    def test_changed_source_exits_with_an_error(self, runner, source):
        runner.invoke(plan_excel, ["-i", str(source)])
        wb = load_workbook(source)
        wb["Catalog"]["A2"] = "changed"
        wb.save(source)
        result = runner.invoke(apply_plan, ["-p", str(default_plan_path(source))])
        assert result.exit_code == 1
        assert "changed since the plan was made" in result.output

    def test_missing_source_is_reported(self, runner, source, tmp_path):
        runner.invoke(plan_excel, ["-i", str(source), "-o", str(tmp_path / "plans" / "p.xlsx")])
        result = runner.invoke(apply_plan, ["-p", str(tmp_path / "plans" / "p.xlsx")])
        assert result.exit_code == 1
        assert "pass it with -i" in result.output

    def test_output_option(self, runner, source, tmp_path):
        runner.invoke(plan_excel, ["-i", str(source)])
        out = tmp_path / "out.xlsx"
        result = runner.invoke(apply_plan, ["-p", str(default_plan_path(source)), "-o", str(out)])
        assert result.exit_code == 0, result.output
        assert out.exists() and default_plan_path(source).exists()
