"""FDP variant after the class questions are answered: FDP range and DASH rules."""

import pytest
from openpyxl import load_workbook
from prep_helpers import CLASS_ANSWERS, answer_classes, cell, find_row, set_cell

from metadata_automation.excel_prep.fdp import derive_fdp_workbook
from metadata_automation.excel_prep.report import Cell


class TestSecondPass:
    def test_class_uri_derived_from_namespace(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(plain_small, first)
        answer_classes(first, class_uris=False)
        second = tmp_path / "second.xlsx"
        derive_fdp_workbook(first, second, namespace="ex")
        ws = load_workbook(second)["classes"]
        assert cell(ws, find_row(ws, "sheet_name", "Catalog"), "class_URI") == "ex:Catalog"

    def test_class_uri_is_not_derived(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(plain_small, first)
        answer_classes(first, class_uris=False)
        second = tmp_path / "second.xlsx"
        report = derive_fdp_workbook(first, second)
        ws = load_workbook(second)["classes"]
        row = find_row(ws, "sheet_name", "Catalog")
        assert not cell(ws, row, "class_URI")
        assert Cell("classes", row, "class_URI") in report.gaps

    def test_independent_class_range_gets_iri_and_relaxed_cardinality(self, second_pass):
        _, wb, _ = second_pass
        catalog, dataset = wb["Catalog"], wb["Dataset"]
        row = find_row(catalog, "Property URI", "dcat:dataset")
        assert cell(catalog, row, "Range") == "dcat:Dataset (IRI)"
        assert cell(catalog, row, "Cardinality") == "0..n"
        row = find_row(dataset, "Property URI", "dcat:distribution")
        assert cell(dataset, row, "Range") == "dcat:Distribution (IRI)"
        assert cell(dataset, row, "Cardinality") == "0..n"
        assert not cell(dataset, row, "SHACL_sh:node")

    def test_changes_name_the_property(self, second_pass):
        report, _, _ = second_pass
        change = next(f for f in report.changes if f.cell.sheet == "Catalog" and f.cell.column == "Cardinality")
        assert change.property_uri == "dcat:dataset"
        assert "(dcat:dataset): '1..n' -> '0..n'" in report.format()
        assert all(f.property_uri is None for f in report.changes if f.cell.sheet == "classes")

    def test_gaps_do_not_name_the_property(self, second_pass):
        report, _, _ = second_pass
        required = report.format().split("Required cells to fill in")[1]
        assert "dct:" not in required and "dcat:" not in required

    def test_dependent_class_range_gets_node_shape(self, second_pass):
        _, wb, _ = second_pass
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        assert cell(ws, row, "Range") == "foaf:Agent"
        assert cell(ws, row, "Cardinality") == "1"
        assert cell(ws, row, "SHACL_sh:node") == "hri:AgentShape"

    def test_differing_node_shape_is_corrected(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        set_cell(ws, row, "SHACL_sh:node", "old:AgentShape")
        wb.save(first)

        second = tmp_path / "second.xlsx"
        report = derive_fdp_workbook(first, second)
        assert any(
            f.cell == Cell("Catalog", row, "SHACL_sh:node")
            and f.message == "'old:AgentShape' -> 'hri:AgentShape' (class of this model)"
            for f in report.changes
        )
        assert cell(load_workbook(second)["Catalog"], row, "SHACL_sh:node") == "hri:AgentShape"

    def test_unset_independent_flag_means_dependent(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(plain_small, first)
        answer_classes(first, {**CLASS_ANSWERS, "Agent": ("foaf:Agent", None)})
        second = tmp_path / "second.xlsx"
        derive_fdp_workbook(first, second)
        ws = load_workbook(second)["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        assert cell(ws, row, "Range") == "foaf:Agent"
        assert cell(ws, row, "SHACL_sh:node") == "hri:AgentShape"

    @pytest.mark.parametrize(
        "sheet, uri, expected",
        [
            ("Catalog", "dct:license", "dct:LicenseDocument (IRI)"),
            ("Agent", "dct:type", "skos:Concept (IRI)"),
            ("Dataset", "dcat:theme", "skos:Concept (IRI)"),
            ("Catalog", "dct:modified", "xsd:dateTime"),
            ("Catalog", "dct:title", "rdfs:Literal"),
            ("Distribution", "dcat:accessURL", "rdfs:Resource (IRI)"),
        ],
    )
    def test_out_of_model_ranges(self, second_pass, sheet, uri, expected):
        _, wb, _ = second_pass
        ws = wb[sheet]
        assert cell(ws, find_row(ws, "Property URI", uri), "Range") == expected

    @pytest.mark.parametrize(
        "uri, viewer, editor",
        [
            ("dcat:dataset", "dash:URIViewer", "dash:URIEditor"),
            ("dct:license", "dash:URIViewer", "dash:URIEditor"),
            ("dct:publisher", "dash:DetailsViewer", "dash:BlankNodeEditor"),
            ("dct:modified", "dash:LiteralViewer", "dash:DateTimePickerEditor"),
            ("dct:description", "dash:LiteralViewer", "dash:TextAreaEditor"),
            ("dct:title", "dash:LiteralViewer", "dash:TextFieldEditor"),
        ],
    )
    def test_empty_dash_cells_are_filled(self, second_pass, uri, viewer, editor):
        _, wb, _ = second_pass
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", uri)
        assert cell(ws, row, "SHACL_dash:viewer") == viewer
        assert cell(ws, row, "SHACL_dash:editor") == editor

    def test_rdfs_class_range_gets_iri(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Distribution"]
        row = find_row(ws, "Property URI", "dcat:accessURL")
        set_cell(ws, row, "Range", "rdfs:Resource")
        wb.save(first)

        second = tmp_path / "second.xlsx"
        derive_fdp_workbook(first, second)
        ws = load_workbook(second)["Distribution"]
        assert cell(ws, row, "Range") == "rdfs:Resource (IRI)"
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"

    def test_generic_resource_gets_uri_viewer(self, second_pass):
        _, wb, _ = second_pass
        ws = wb["Distribution"]
        row = find_row(ws, "Property URI", "dcat:accessURL")
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"

    def test_is_idempotent(self, second_pass, tmp_path):
        _, _, path = second_pass
        report = derive_fdp_workbook(path, tmp_path / "third.xlsx")
        assert report.changes == []

    def test_no_title_issue_when_independent_classes_have_one(self, second_pass):
        report, _, _ = second_pass
        assert not any("dct:title" in f.message for f in report.issues)
