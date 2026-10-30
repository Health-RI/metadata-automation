"""Small pure helpers."""

import pytest
from openpyxl import Workbook

from metadata_automation.excel_prep.dash import LABEL, URI, expected_dash, has_vocab
from metadata_automation.excel_prep.fdp import relax_cardinality
from metadata_automation.excel_prep.report import compress_rows
from metadata_automation.excel_prep.workbook import is_literal


@pytest.mark.parametrize(
    "given, expected",
    [("1", "0..1"), ("1..1", "0..1"), ("1..n", "0..n"), ("2..5", "0..5"), ("0..n", "0..n"), ("0..1", "0..1")],
)
def test_relax_cardinality(given, expected):
    assert relax_cardinality(given) == expected


def test_relax_cardinality_unparsable():
    assert relax_cardinality("one or more") is None


def test_compress_rows():
    assert compress_rows([7, 2, 3, 4, 9, 10]) == "2-4, 7, 9-10"


@pytest.mark.parametrize(
    "range_, expected",
    [("rdfs:Literal", True), ("xsd:string", True), ("rdfs:Resource", False), ("rdfs:Class", False)],
)
def test_is_literal(range_, expected):
    assert is_literal(range_) is expected


ACCESS_RIGHT = "http://publications.europa.eu/resource/authority/access-right"


@pytest.mark.parametrize(
    "value, hyperlink, expected",
    [
        ("", None, False),
        ("NA", None, False),
        ("GeoNames", None, False),
        ("http://example.org/vocab", None, False),
        (ACCESS_RIGHT, None, True),
        ("Access rights", ACCESS_RIGHT, True),
    ],
)
def test_has_vocab(value, hyperlink, expected):
    # Arrange
    ws = Workbook().active
    ws.cell(row=2, column=1, value=value or None).hyperlink = hyperlink

    # Act
    result = has_vocab(ws, 2, 1)

    # Assert
    assert result is expected


def test_has_vocab_without_column():
    assert has_vocab(Workbook().active, 2, None) is False


@pytest.mark.parametrize(
    "range_, vocab, expected",
    [
        ("dct:LicenseDocument (IRI)", True, LABEL),
        ("dct:LicenseDocument (IRI)", False, URI),
        ("rdfs:Resource (IRI)", True, URI),
        ("dcat:Dataset (IRI)", True, URI),
    ],
)
def test_expected_dash_label_viewer_needs_vocab(range_, vocab, expected):
    assert expected_dash(range_, "", False, {"dcat:Dataset"}, vocab) == expected
