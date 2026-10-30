"""Planning the technical columns: the source is untouched, the plan lists what would change."""

import shutil

import pytest
from openpyxl import Workbook, load_workbook
from prep_helpers import CLASS_ANSWERS, answer_classes, fill_reported_gaps, sha

from metadata_automation.cli import plan_excel
from metadata_automation.excel_prep.plan import (
    ADD_COLUMN,
    ADD_ROW,
    ADD_SHEET,
    RENAME_SHEET,
    SET_CELL,
    Edit,
    compute_plan,
    default_plan_path,
    read_plan,
    row_key,
    write_plan,
)
from metadata_automation.excel_prep.preparer import prepare_workbook
from metadata_automation.excel_prep.workbook import headers


@pytest.fixture
def source(plain_small, tmp_path):
    path = tmp_path / "model.xlsx"
    shutil.copy(plain_small, path)
    return path


@pytest.fixture
def prepared(source):
    """The source after the class questions are answered and prepared, so most rules apply."""
    prepare_workbook(source)
    answer_classes(source, {k: v for k, v in CLASS_ANSWERS.items() if k != "Resource"})
    prepare_workbook(source)
    return source


class TestComputePlan:
    def test_source_is_not_modified(self, source):
        before = sha(source)
        compute_plan(source)
        assert sha(source) == before

    def test_summary_identifies_the_source(self, source):
        plan, _ = compute_plan(source)
        assert plan.source_name == "model.xlsx"
        assert plan.source_hash == sha(source)

    def test_structural_edits_for_a_plain_workbook(self, source):
        plan, _ = compute_plan(source)
        kinds = {e.kind for e in plan.edits}
        assert {ADD_SHEET, ADD_COLUMN, ADD_ROW} <= kinds
        assert {e.sheet for e in plan.edits if e.kind == ADD_SHEET} == {"prefixes", "classes"}
        assert {e.key for e in plan.edits if e.kind == ADD_ROW and e.sheet == "classes"} == {
            "Catalog",
            "Dataset",
            "Agent",
            "Distribution",
        }
        added = {e.column for e in plan.edits if e.kind == ADD_COLUMN and e.sheet == "Catalog"}
        assert {"SHACL_sh:node", "SeMPyRO_range"} <= added

    def test_added_row_carries_its_key_value(self, source):
        plan, _ = compute_plan(source)
        row = next(e for e in plan.edits if e.kind == ADD_ROW and e.key == "Dataset")
        assert (row.column, row.new) == ("sheet_name", "Dataset")

    def test_cell_edits_name_the_property_and_the_reason(self, prepared):
        wb = load_workbook(prepared)
        ws = wb["Catalog"]
        cols = headers(ws)
        row = next(
            r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=cols["Property URI"]).value == "dct:publisher"
        )
        ws.cell(row=row, column=cols["SHACL_sh:node"], value="old:AgentShape")
        wb.save(prepared)

        plan, _ = compute_plan(prepared)
        edit = next(e for e in plan.edits if e.kind == SET_CELL)
        assert (edit.sheet, edit.key, edit.column) == ("Catalog", "dct:publisher", "SHACL_sh:node")
        assert (edit.old, edit.new) == ("old:AgentShape", "hri:AgentShape")
        assert edit.reason == "class of this model"
        assert edit.apply is True

    def test_prepared_workbook_has_no_edits(self, prepared):
        plan, _ = compute_plan(prepared)
        assert plan.edits == []

    def test_gaps_are_required_cells_only_and_keyed_by_content(self, prepared):
        plan, _ = compute_plan(prepared)
        assert not plan.is_empty
        columns = {(g.sheet, g.column) for g in plan.gaps}
        assert ("Catalog", "SeMPyRO_rdf_term") in columns
        assert ("classes", "SeMPyRO_annotations_IRI") in columns
        assert all(g.key and g.value == "" for g in plan.gaps)
        assert not any(g.column == "SHACL_pattern" for g in plan.gaps)

    def test_gaps_in_ignored_columns_are_counted_not_listed(self, prepared):
        everything, _ = compute_plan(prepared)
        plan, _ = compute_plan(prepared, ignore_gaps=["SeMPyRO_*"])
        assert not any(g.column.startswith("SeMPyRO_") for g in plan.gaps)
        assert plan.ignored_gaps == len(everything.gaps) - len(plan.gaps) > 0
        assert plan.ignored_columns == "SeMPyRO_*"

    def test_ignored_gaps_are_kept_in_the_plan_file(self, prepared, tmp_path):
        plan, _ = compute_plan(prepared, ignore_gaps=["SeMPyRO_*", "SHACL_pattern"])
        write_plan(plan, tmp_path / "p.xlsx")
        loaded = read_plan(tmp_path / "p.xlsx")
        assert (loaded.ignored_columns, loaded.ignored_gaps) == ("SeMPyRO_*, SHACL_pattern", plan.ignored_gaps)

    def test_fully_prepared_and_filled_workbook_gives_an_empty_plan(self, prepared):
        fill_reported_gaps(prepared)
        prepare_workbook(prepared)
        plan, _ = compute_plan(prepared)
        assert plan.is_empty, (plan.edits[:3], plan.gaps[:3])

    def test_duplicate_keys_get_an_occurrence_number(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Thing"
        ws.append(["Property label", "Property URI", "Range", "Cardinality"])
        ws.append(["a", "ex:p", "xsd:string", "1"])
        ws.append(["b", "ex:p", "xsd:string", "1"])
        path = tmp_path / "dup.xlsx"
        wb.save(path)

        plan, _ = compute_plan(path)
        gaps = [g for g in plan.gaps if g.sheet == "Thing" and g.column == "SeMPyRO_rdf_term"]
        assert [(g.key, g.occurrence) for g in gaps] == [("ex:p", 1), ("ex:p", 2)]

    def test_case_variant_of_a_general_sheet_is_a_rename(self, source):
        wb = load_workbook(source)
        wb.create_sheet("Classes", 0)
        wb["Classes"].append(["sheet_name"])
        wb.save(source)

        plan, _ = compute_plan(source)
        rename = next(e for e in plan.edits if e.kind == RENAME_SHEET)
        assert (rename.old, rename.new) == ("Classes", "classes")

    def test_row_key_falls_back_to_label_and_row_number(self):
        ws = Workbook().active
        ws.append(["Property label", "Property URI"])
        ws.append(["only a label", None])
        ws.append([None, None])
        cols = headers(ws)
        assert row_key(ws, 2, cols) == "label: only a label"
        assert row_key(ws, 3, cols) == "row 3"


class TestPlanFile:
    def test_round_trip(self, source, tmp_path):
        plan, _ = compute_plan(source, namespace="ex")
        path = tmp_path / "plan.xlsx"
        write_plan(plan, path)

        loaded = read_plan(path)
        assert loaded.source_name == plan.source_name
        assert loaded.source_hash == plan.source_hash
        assert loaded.namespace == "ex"
        assert loaded.edits == plan.edits
        assert loaded.gaps == plan.gaps
        assert loaded.issues == plan.issues

    def test_edits_by_a_person_are_read_back(self, source, tmp_path):
        plan, _ = compute_plan(source)
        path = tmp_path / "plan.xlsx"
        write_plan(plan, path)
        wb = load_workbook(path)
        edits, gaps = wb["Edits"], wb["Gaps"]
        edits.cell(row=2, column=1, value="no")
        gaps.cell(row=2, column=5, value="hri:Catalog")
        wb.save(path)

        loaded = read_plan(path)
        assert loaded.edits[0].apply is False
        assert all(e.apply for e in loaded.edits[1:])
        assert loaded.gaps[0].value == "hri:Catalog"

    def test_existing_plan_is_not_overwritten(self, source, tmp_path):
        plan, _ = compute_plan(source)
        path = tmp_path / "plan.xlsx"
        write_plan(plan, path)
        with pytest.raises(ValueError, match="--overwrite"):
            write_plan(plan, path)
        write_plan(plan, path, overwrite=True)

    def test_a_workbook_that_is_not_a_plan_is_rejected(self, source):
        with pytest.raises(ValueError, match="not a plan"):
            read_plan(source)

    def test_default_path_is_next_to_the_source(self, tmp_path):
        assert default_plan_path(tmp_path / "Health-RI.xlsx") == tmp_path / "Health-RI-plan.xlsx"

    def test_edit_defaults(self):
        assert Edit(SET_CELL, "s").apply is True


class TestPlanCommand:
    def test_writes_the_plan_next_to_the_source(self, runner, source):
        before = sha(source)
        result = runner.invoke(plan_excel, ["-i", str(source)])
        assert result.exit_code == 0, result.output
        assert sha(source) == before
        assert default_plan_path(source).exists()
        assert "Plan written to" in result.output
        assert "Required cells to fill in" in result.output

    def test_refuses_to_replace_a_plan_without_overwrite(self, runner, source):
        assert runner.invoke(plan_excel, ["-i", str(source)]).exit_code == 0
        result = runner.invoke(plan_excel, ["-i", str(source)])
        assert result.exit_code == 1
        assert "--overwrite" in result.output
        assert runner.invoke(plan_excel, ["-i", str(source), "--overwrite"]).exit_code == 0

    def test_custom_output_path(self, runner, source, tmp_path):
        out = tmp_path / "elsewhere" / "p.xlsx"
        result = runner.invoke(plan_excel, ["-i", str(source), "-o", str(out)])
        assert result.exit_code == 0, result.output
        assert out.exists() and not default_plan_path(source).exists()
