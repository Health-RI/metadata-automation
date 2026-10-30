"""Helpers shared by the Excel preparation tests."""

import hashlib

from openpyxl import load_workbook

from metadata_automation.excel_prep.workbook import class_sheet_names, data_rows, headers, property_rows, text

# Answers a modeler gives after the first pass: target class and whether it is independent.
CLASS_ANSWERS = {
    "Resource": ("dcat:Resource", False),
    "Catalog": ("dcat:Catalog", True),
    "Dataset": ("dcat:Dataset", True),
    "Agent": ("foaf:Agent", False),
    "Distribution": ("dcat:Distribution", True),
}
NAMESPACES = {
    "hri": "https://w3id.org/health-ri/metadata-vocabulary#",
    "dct": "http://purl.org/dc/terms/",
    "dcat": "http://www.w3.org/ns/dcat#",
    "foaf": "http://xmlns.com/foaf/0.1/",
    "rdfs": "http://www.w3.org/2000/01/rdf-schema#",
    "xsd": "http://www.w3.org/2001/XMLSchema#",
    "skos": "http://www.w3.org/2004/02/skos/core#",
    "dash": "http://datashapes.org/dash#",
}
RDF_NAMESPACES = {"dct": "DCTERMS", "dcat": "DCAT", "foaf": "FOAF"}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def find_row(ws, column, value):
    cols = headers(ws)
    return next(r for r in range(2, ws.max_row + 1) if text(ws, r, cols[column]) == value)


def cell(ws, row, column):
    return text(ws, row, headers(ws)[column])


def set_cell(ws, row, column, value):
    ws.cell(row=row, column=headers(ws)[column], value=value)


def answer_classes(path, answers=CLASS_ANSWERS, class_uris=True):
    """Fill in the target, independent flag and (optionally) class_URI, as the modeler would."""
    wb = load_workbook(path)
    ws = wb["classes"]
    for name, (target, independent) in answers.items():
        row = find_row(ws, "sheet_name", name)
        set_cell(ws, row, "SHACL_target_ontology_name", target)
        set_cell(ws, row, "SHACL_independent_class", independent)
        if class_uris:
            set_cell(ws, row, "class_URI", f"hri:{target.split(':')[1]}")
    wb.save(path)


def rename_sheet(wb, old, new):
    """Rename through a free name; openpyxl counts the sheet's own name as a clash in any case."""
    wb[old].title = f"{new}__tmp"
    wb[f"{new}__tmp"].title = new


def fill_reported_gaps(path):
    """Fill every required gap with a valid value, as the modeler would after reading the report."""
    wb = load_workbook(path)
    ws = wb["prefixes"]
    for row in data_rows(ws, headers(ws), ("prefix",)):
        set_cell(ws, row, "namespace", NAMESPACES[cell(ws, row, "prefix")])

    ws = wb["classes"]
    for row in data_rows(ws, headers(ws), ("sheet_name",)):
        set_cell(ws, row, "SeMPyRO_annotations_ontology", "https://www.w3.org/TR/vocab-dcat-3/")
        set_cell(ws, row, "SeMPyRO_annotations_IRI", f"https://example.org/{cell(ws, row, 'sheet_name')}")
        set_cell(ws, row, "SeMPyRO_add_rdf_model", True)

    for name in class_sheet_names(wb):
        ws = wb[name]
        for row in property_rows(ws, headers(ws)):
            prefix, local = cell(ws, row, "Property URI").split(":")
            literal = cell(ws, row, "Range").startswith(("rdfs:Literal", "xsd:"))
            set_cell(ws, row, "SeMPyRO_rdf_term", f"{RDF_NAMESPACES[prefix]}.{local}")
            set_cell(ws, row, "SeMPyRO_rdf_type", "rdfs_literal" if literal else "uri")
            set_cell(ws, row, "SeMPyRO_range", "LiteralField, str" if literal else "AnyHttpUrl")
    wb.save(path)
