"""Writing the result without losing the source or an existing output."""

import os
import re
import shutil
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from prep_helpers import sha

from metadata_automation.excel_prep.drawings import _drawing_target, _read, _sheet_parts, _write, drawing_sheets
from metadata_automation.excel_prep.preparer import prepare_workbook
from metadata_automation.excel_prep.report import Cell

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d4944415478da63f8ffff3f0005fe02fe0dc8f5a30000000049454e44ae426082"
)
DRAWING = (
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="{rel}">'
    "<xdr:oneCellAnchor><xdr:from><xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row>"
    '<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="2000000" cy="500000"/>'
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="TextBox 1"/><xdr:cNvSpPr txBox="1"/>'
    '</xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2000000" cy="500000"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr><xdr:txBody><a:bodyPr/><a:lstStyle/>'
    "<a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:oneCellAnchor>"
    "<xdr:oneCellAnchor><xdr:from><xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row>"
    '<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="500000" cy="500000"/>'
    '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="3" name="Picture 2"/><xdr:cNvPicPr/></xdr:nvPicPr>'
    '<xdr:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
    '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="500000" cy="500000"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
    "</xdr:wsDr>"
)


def rels(*relations):
    body = "".join(
        f'<Relationship Id="rId{i}" Type="{REL}/{t}" Target="{target}"/>' for i, (t, target) in enumerate(relations, 1)
    )
    return f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'


def add_drawings(path, sheet_indices):
    """Give each sheet a text box and a picture; the pictures share one image."""
    tmp = path.with_suffix(".tmp")
    with ZipFile(path) as src, ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                overrides = "".join(
                    f'<Override PartName="/xl/drawings/drawing{n}.xml"'
                    ' ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                    for n in sheet_indices
                )
                data = data.replace(
                    b"</Types>", f'<Default Extension="png" ContentType="image/png"/>{overrides}</Types>'.encode()
                )
            dst.writestr(item, data)
        for n in sheet_indices:
            dst.writestr(f"xl/worksheets/_rels/sheet{n}.xml.rels", rels(("drawing", f"../drawings/drawing{n}.xml")))
            dst.writestr(f"xl/drawings/drawing{n}.xml", DRAWING.format(rel=REL, text=f"note {n}"))
            dst.writestr(f"xl/drawings/_rels/drawing{n}.xml.rels", rels(("image", "../media/image1.png")))
        dst.writestr("xl/media/image1.png", PNG)
    os.replace(tmp, path)


@pytest.fixture
def plain_copy(plain_small, tmp_path):
    path = tmp_path / "plain.xlsx"
    shutil.copy(plain_small, path)
    return path


def add_chart(path, sheet, text_box=False):
    """Add an openpyxl chart, optionally with a text box in the same drawing."""
    wb = load_workbook(path)
    chart = BarChart()
    chart.add_data(Reference(wb[sheet], min_col=1, min_row=1, max_row=2))
    wb[sheet].add_chart(chart, "H2")
    wb.save(path)
    if text_box:
        files = _read(path)
        drawing = next(n for n in files if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
        anchor = re.search(rb"<xdr:oneCellAnchor>.*?</xdr:oneCellAnchor>", DRAWING.format(rel=REL, text="x").encode())
        # openpyxl writes the drawing namespace as default, without the xdr prefix.
        files[drawing] = files[drawing].replace(b"</wsDr>", anchor.group(0).replace(b"xdr:", b"") + b"</wsDr>")
        _write(path, files)


def drawing_texts(path):
    """Sheet name to the text box text of its drawing."""
    files = _read(path)
    texts = {}
    for name, part in _sheet_parts(files).items():
        drawing = _drawing_target(files, part)
        if drawing and (text := re.search(rb"<a:t>(.*?)</a:t>", files[drawing])):
            texts[name] = text.group(1).decode()
    return texts


def test_existing_output_is_refused(plain_small, tmp_path):
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"keep")
    with pytest.raises(ValueError, match="--overwrite"):
        prepare_workbook(plain_small, out)
    assert out.read_bytes() == b"keep"


def test_overwrite_replaces_existing_output(plain_small, tmp_path):
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"old")
    prepare_workbook(plain_small, out, overwrite=True)
    assert "classes" in load_workbook(out).sheetnames


def test_same_file_by_other_path_is_refused(plain_copy, tmp_path):
    link = tmp_path / "link.xlsx"
    os.link(plain_copy, link)
    before = sha(plain_copy)
    with pytest.raises(ValueError, match="must differ"):
        prepare_workbook(plain_copy, link, overwrite=True)
    assert sha(plain_copy) == before


def test_failed_save_leaves_source_intact(plain_copy, monkeypatch):
    def fail(self, filename):
        open(filename, "wb").write(b"partial")
        raise OSError("disk full")

    before = sha(plain_copy)
    monkeypatch.setattr(Workbook, "save", fail)
    with pytest.raises(OSError):
        prepare_workbook(plain_copy, init=True)
    assert sha(plain_copy) == before
    assert {p.name for p in plain_copy.parent.iterdir()} == {plain_copy.name}


def test_init_keeps_file_mode(plain_copy):
    plain_copy.chmod(0o640)
    prepare_workbook(plain_copy, init=True)
    assert plain_copy.stat().st_mode & 0o777 == 0o640


def test_new_output_gets_default_mode(plain_small, tmp_path):
    out = tmp_path / "out.xlsx"
    prepare_workbook(plain_small, out)
    umask = os.umask(0)
    os.umask(umask)
    assert out.stat().st_mode & 0o777 == 0o666 & ~umask


def test_drawing_sheets(plain_copy):
    add_drawings(plain_copy, [1, 5])
    assert drawing_sheets(plain_copy) == ["Info", "Distribution"]


def test_drawings_are_kept_by_sheet_name(plain_copy, tmp_path):
    add_drawings(plain_copy, [1, 5])
    out = tmp_path / "out.xlsx"
    report = prepare_workbook(plain_copy, out)

    assert drawing_texts(out) == {"Info": "note 1", "Distribution": "note 5"}
    files = _read(out)
    assert files["xl/media/image1.png"] == PNG
    assert sum(n.startswith("xl/media/") for n in files) == 1
    assert b'PartName="/xl/drawings/drawing5.xml"' in files["[Content_Types].xml"]
    assert not any("drawings" in f.message for f in report.issues)
    load_workbook(out)


def test_drawings_follow_a_renamed_general_sheet(plain_copy, tmp_path):
    add_drawings(plain_copy, [1])
    files = _read(plain_copy)
    files["xl/workbook.xml"] = files["xl/workbook.xml"].replace(b'name="Info"', b'name="info"')
    _write(plain_copy, files)

    out = tmp_path / "out.xlsx"
    prepare_workbook(plain_copy, out)
    assert drawing_texts(out) == {"Info": "note 1"}


def test_init_keeps_drawings_and_backs_up_source(plain_copy):
    add_drawings(plain_copy, [1])
    before = sha(plain_copy)
    report = prepare_workbook(plain_copy, init=True)
    assert sha(plain_copy.with_name("plain.xlsx.bak")) == before
    assert any(f.cell == Cell("plain.xlsx") and "plain.xlsx.bak" in f.message for f in report.changes)
    assert drawing_texts(plain_copy) == {"Info": "note 1"}


def test_full_run_writes_no_backup(plain_copy, tmp_path):
    add_drawings(plain_copy, [1])
    prepare_workbook(plain_copy, tmp_path / "out.xlsx")
    assert not plain_copy.with_name("plain.xlsx.bak").exists()


def test_chart_kept_by_openpyxl_is_not_reported(plain_copy, tmp_path):
    add_chart(plain_copy, "Info")
    add_drawings(plain_copy, [5])
    report = prepare_workbook(plain_copy, tmp_path / "out.xlsx")
    assert drawing_texts(tmp_path / "out.xlsx")["Distribution"] == "note 5"
    assert not any("drawings" in f.message for f in report.issues)


def test_shapes_next_to_chart_are_reported(plain_copy, tmp_path):
    add_chart(plain_copy, "Info", text_box=True)
    report = prepare_workbook(plain_copy, tmp_path / "out.xlsx")
    assert any(f.cell == Cell("Info") and "could not be kept" in f.message for f in report.issues)


def test_rich_text_is_kept(plain_copy):
    wb = load_workbook(plain_copy)
    rich = CellRichText(["plain ", TextBlock(InlineFont(b=True), "bold")])
    wb["Info"]["A10"] = rich
    wb.save(plain_copy)

    prepare_workbook(plain_copy, init=True)
    value = load_workbook(plain_copy, rich_text=True)["Info"]["A10"].value
    assert isinstance(value, CellRichText)
    assert str(value) == "plain bold"
