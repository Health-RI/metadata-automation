"""Init run: in-place update without the FDP rules."""

import shutil

import pytest
from openpyxl import load_workbook
from prep_helpers import CLASS_ANSWERS, answer_classes, cell, find_row, set_cell, sha

from metadata_automation.excel_prep.columns import CLASS_SHEET_COLUMNS, VOCAB_COLUMN
from metadata_automation.excel_prep.fdp import derive_fdp_workbook
from metadata_automation.excel_prep.preparer import prepare_workbook
from metadata_automation.excel_prep.report import Cell
from metadata_automation.excel_prep.workbook import class_sheet_names, headers


class TestInit:
    @pytest.fixture
    def plain_copy(self, plain_small, tmp_path):
        path = tmp_path / "plain.xlsx"
        shutil.copy(plain_small, path)
        return path

    @pytest.fixture
    def answered(self, plain_copy):
        """An init output with the class questions answered, so every rule could apply."""
        prepare_workbook(plain_copy)
        answer_classes(plain_copy, {k: v for k, v in CLASS_ANSWERS.items() if k != "Resource"})
        return plain_copy

    def test_updates_in_place_without_fdp_sheets(self, plain_copy):
        before = sha(plain_copy)
        report = prepare_workbook(plain_copy)
        assert sha(plain_copy) != before
        wb = load_workbook(plain_copy)
        assert wb.sheetnames[:3] == ["Info", "prefixes", "classes"]
        assert "Resource" not in wb.sheetnames
        for name in class_sheet_names(wb):
            assert set(CLASS_SHEET_COLUMNS) <= set(headers(wb[name]))
        assert any(c.cell.sheet == "prefixes" for c in report.changes)
        assert report.prefix_issues

    def test_output_path_writes_a_new_file_without_fdp_rules(self, plain_copy, tmp_path):
        out = tmp_path / "out.xlsx"
        before = plain_copy.read_bytes()
        prepare_workbook(plain_copy, out)
        assert plain_copy.read_bytes() == before
        assert "Resource" not in load_workbook(out).sheetnames

    def test_model_range_gets_node_shape_only(self, answered):
        report = prepare_workbook(answered)
        ws = load_workbook(answered)["Catalog"]
        row = find_row(ws, "Property URI", "dcat:dataset")
        assert cell(ws, row, "Range") == "dcat:Dataset"
        assert cell(ws, row, "Cardinality") == "1..n"
        assert cell(ws, row, "SHACL_sh:node") == "hri:DatasetShape"
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:DetailsViewer"
        untouched = {Cell("Catalog", row, "Range"), Cell("Catalog", row, "Cardinality")}
        assert not any(f.cell in untouched for f in report.changes)

    def test_out_of_model_range_gets_iri_and_uri_viewer(self, answered):
        report = prepare_workbook(answered)
        ws = load_workbook(answered)["Catalog"]
        row = find_row(ws, "Property URI", "dct:license")
        assert cell(ws, row, "Range") == "dct:LicenseDocument (IRI)"
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"
        assert cell(ws, row, "SHACL_dash:editor") == "dash:URIEditor"
        assert any(
            f.cell == Cell("Catalog", row, "Range")
            and f.message == "'dct:LicenseDocument' -> 'dct:LicenseDocument (IRI)' (range is not a class of this model)"
            for f in report.changes
        )

    def test_out_of_model_range_with_vocab_gets_label_viewer(self, answered):
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:license")
        set_cell(ws, row, VOCAB_COLUMN, "http://publications.europa.eu/resource/authority/access-right")
        wb.save(answered)
        prepare_workbook(answered)
        ws = load_workbook(answered)["Catalog"]
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:LabelViewer"
        assert cell(ws, row, "SHACL_dash:editor") == "dash:URIEditor"

    def test_out_of_model_range_with_unknown_vocab_gets_uri_viewer(self, answered):
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:license")
        set_cell(ws, row, VOCAB_COLUMN, "GeoNames")
        wb.save(answered)
        prepare_workbook(answered)
        ws = load_workbook(answered)["Catalog"]
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"

    def test_missing_vocab_column_is_reported_once(self, answered):
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        ws.delete_cols(headers(ws)[VOCAB_COLUMN])
        wb.save(answered)
        report = prepare_workbook(answered)
        missing = [f for f in report.issues if f.cell == Cell("Catalog", None, VOCAB_COLUMN)]
        assert len(missing) == 1

    def test_full_run_references_independent_class_by_iri(self, answered, tmp_path):
        prepare_workbook(answered)
        first = tmp_path / "first.xlsx"
        derive_fdp_workbook(answered, first)
        answer_classes(first)
        out = tmp_path / "fdp.xlsx"
        report = derive_fdp_workbook(first, out)
        ws = load_workbook(out)["Catalog"]
        row = find_row(ws, "Property URI", "dcat:dataset")
        assert cell(ws, row, "Range") == "dcat:Dataset (IRI)"
        assert not cell(ws, row, "SHACL_sh:node")
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"
        assert cell(ws, row, "SHACL_dash:editor") == "dash:URIEditor"
        assert "cleared 'hri:DatasetShape' ('Dataset' is an independent class)" in [f.message for f in report.changes]
        assert not any(f.cell.row == row and f.cell.sheet == "Catalog" for f in report.issues)

    def test_no_title_check(self, answered):
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        ws.delete_rows(find_row(ws, "Property URI", "dct:title"))
        wb.save(answered)
        report = prepare_workbook(answered)
        assert not any("dct:title" in f.message for f in report.issues)
