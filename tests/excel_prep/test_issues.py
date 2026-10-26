"""Problems that are reported instead of fixed."""

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from prep_helpers import CLASS_ANSWERS, answer_classes, cell, find_row, set_cell, sha

from metadata_automation.excel_prep.preparer import prepare_workbook
from metadata_automation.excel_prep.report import Cell
from metadata_automation.excel_prep.workbook import headers


class TestIssues:
    def test_dependent_range_with_iri_is_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Catalog"]
        set_cell(ws, find_row(ws, "Property URI", "dct:publisher"), "Range", "foaf:Agent (IRI)")
        wb.save(first)

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        issue = next(f for f in report.issues if f.cell.sheet == "Catalog" and "dependent class" in f.message)
        assert issue.property_uri == "dct:publisher"
        assert "Catalog / Range / row" in str(issue) and "(dct:publisher):" in str(issue)

    def test_differing_dash_value_is_reported_not_overwritten(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        set_cell(ws, row, "SHACL_dash:viewer", "dash:URIViewer")
        wb.save(first)

        second = tmp_path / "second.xlsx"
        report = prepare_workbook(first, second)
        assert any(
            f.cell == Cell("Catalog", row, "SHACL_dash:viewer")
            and f.message == "expected 'dash:DetailsViewer', found 'dash:URIViewer'"
            for f in report.issues
        )
        ws = load_workbook(second)["Catalog"]
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"
        assert cell(ws, row, "SHACL_dash:editor") == "dash:BlankNodeEditor"

    def test_no_dash_values_before_targets_are_known(self, plain_small, tmp_path):
        out = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, out)
        ws = load_workbook(out)["Catalog"]
        assert not cell(ws, find_row(ws, "Property URI", "dct:title"), "SHACL_dash:viewer")

    def test_independent_class_without_title_is_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first, {**CLASS_ANSWERS, "Agent": ("foaf:Agent", True)})

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        assert any(f.cell.sheet == "Agent" and "no dct:title" in f.message for f in report.issues)

    def test_invalid_flag_and_unknown_sheet_are_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        wb = load_workbook(first)
        ws = wb["classes"]
        set_cell(ws, find_row(ws, "sheet_name", "Agent"), "SHACL_independent_class", "maybe")
        ws.cell(row=ws.max_row + 1, column=1, value="Ghost")
        wb.save(first)

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        messages = [f.message for f in report.issues]
        assert "expected TRUE or FALSE, got 'maybe'" in messages
        assert "sheet 'Ghost' does not exist" in messages

    def test_unheaded_data_is_reported_and_kept(self, plain_small, tmp_path):
        src = tmp_path / "src.xlsx"
        wb = load_workbook(plain_small)
        ws = wb["Catalog"]
        notes_col = ws.max_column + 2
        ws.cell(row=3, column=notes_col, value="modeler note")
        wb.save(src)

        out = tmp_path / "first.xlsx"
        report = prepare_workbook(src, out)
        letter = get_column_letter(notes_col)
        assert any(
            f.cell == Cell("Catalog", None, f"column {letter}") and f.message == "data in column without header"
            for f in report.issues
        )
        ws = load_workbook(out)["Catalog"]
        assert ws.cell(row=3, column=notes_col).value == "modeler note"
        assert all(col > notes_col for name, col in headers(ws).items() if name.startswith("SHACL_"))

    def test_untargeted_sheets_are_named(self, plain_small, tmp_path):
        report = prepare_workbook(plain_small, tmp_path / "first.xlsx")
        issue = next(f for f in report.issues if "range rules skipped" in f.message)
        assert all(name in issue.message for name in CLASS_ANSWERS)

    def test_duplicate_target_blocks_the_run(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first, {**CLASS_ANSWERS, "Dataset": ("dcat:Catalog", True)})
        ws = load_workbook(first)["classes"]
        rows = ", ".join(str(find_row(ws, "sheet_name", n)) for n in ("Catalog", "Dataset"))
        before = sha(first)

        with pytest.raises(ValueError, match=f"SHACL_target_ontology_name 'dcat:Catalog' in rows {rows}"):
            prepare_workbook(first, init=True)
        assert sha(first) == before
        assert not first.with_name("first.xlsx.bak").exists()

    def test_duplicate_sheet_name_blocks_the_run(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        wb = load_workbook(first)
        ws = wb["classes"]
        set_cell(ws, find_row(ws, "sheet_name", "Agent"), "sheet_name", "Dataset")
        wb.save(first)

        out = tmp_path / "second.xlsx"
        with pytest.raises(ValueError, match="sheet_name 'Dataset' in rows"):
            prepare_workbook(first, out)
        assert not out.exists()

    def test_property_without_uri_is_a_gap(self, plain_small, tmp_path):
        src = tmp_path / "src.xlsx"
        wb = load_workbook(plain_small)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:title")
        ws.cell(row=row, column=headers(ws)["Property URI"]).value = None
        wb.save(src)

        report = prepare_workbook(src, tmp_path / "out.xlsx")
        assert Cell("Catalog", row, "Property URI") in report.gaps
