"""Full run on a plain workbook: structure is added, nothing is guessed."""

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from prep_helpers import cell, sha

from metadata_automation.excel_prep.columns import CLASS_SHEET_COLUMNS, CLASSES_COLUMNS
from metadata_automation.excel_prep.fdp import derive_fdp_workbook
from metadata_automation.excel_prep.workbook import class_sheet_names, data_rows, headers, text


class TestFirstPass:
    @pytest.fixture
    def prepared(self, plain_small, tmp_path):
        out = tmp_path / "prepared.xlsx"
        report = derive_fdp_workbook(plain_small, out)
        return report, load_workbook(out)

    def test_source_is_untouched(self, plain_small, tmp_path):
        before = sha(plain_small)
        derive_fdp_workbook(plain_small, tmp_path / "out.xlsx")
        assert sha(plain_small) == before

    def test_refuses_to_overwrite_source(self, plain_small):
        with pytest.raises(ValueError, match="differ"):
            derive_fdp_workbook(plain_small, plain_small)

    def test_general_sheets_added_after_info(self, prepared):
        _, wb = prepared
        assert wb.sheetnames[:4] == ["Info", "prefixes", "classes", "Resource"]
        assert wb["Info"]["A2"].value == "Plain test model"

    def test_classes_sheet_has_row_per_class_sheet(self, prepared):
        _, wb = prepared
        ws = wb["classes"]
        cols = headers(ws)
        assert list(cols) == CLASSES_COLUMNS
        names = [text(ws, r, cols["sheet_name"]) for r in data_rows(ws, cols, ("sheet_name",))]
        assert names == class_sheet_names(wb) == ["Resource", "Catalog", "Dataset", "Agent", "Distribution"]

    def test_class_sheets_have_pipeline_columns(self, prepared):
        _, wb = prepared
        for name in class_sheet_names(wb):
            assert set(CLASS_SHEET_COLUMNS) <= set(headers(wb[name]))

    def test_added_columns_copy_first_column_format(self, plain_small, tmp_path):
        styled = tmp_path / "styled.xlsx"
        wb = load_workbook(plain_small)
        ws = wb["Dataset"]
        ws["A1"].fill = PatternFill("solid", fgColor="FFC000")
        ws["A1"].font = Font(bold=True, color="FF0000")
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 31
        wb.save(styled)

        out = tmp_path / "prepared.xlsx"
        derive_fdp_workbook(styled, out)
        ws = load_workbook(out)["Dataset"]
        for name in CLASS_SHEET_COLUMNS:
            col = headers(ws)[name]
            header, data = ws.cell(row=1, column=col), ws.cell(row=2, column=col)
            assert header.fill.fgColor.rgb == "00FFC000"
            assert header.font.b and header.font.color.rgb == "00FF0000"
            assert data.alignment.wrap_text
            assert ws.column_dimensions[get_column_letter(col)].width == 31

    def test_prefixes_sheet_lists_used_prefixes_without_namespace(self, prepared):
        report, wb = prepared
        ws = wb["prefixes"]
        cols = headers(ws)
        prefixes = {text(ws, r, cols["prefix"]) for r in data_rows(ws, cols, ("prefix",))}
        assert prefixes == {"dcat", "dct", "foaf", "rdfs", "skos", "xsd"}
        assert all(not text(ws, r, cols["namespace"]) for r in data_rows(ws, cols, ("prefix",)))
        assert any("'dct' has no registered namespace" in f.message for f in report.prefix_issues)

    def test_nothing_is_guessed(self, prepared):
        report, wb = prepared
        ws = wb["classes"]
        for row in data_rows(ws, headers(ws), ("sheet_name",)):
            for column in ("class_URI", "SHACL_target_ontology_name", "SHACL_independent_class"):
                assert not cell(ws, row, column)
        assert cell(wb["Catalog"], 3, "Range") == "dcat:Dataset"
        assert any("range rules skipped" in f.message for f in report.issues)

    def test_report_lists_gaps(self, prepared):
        report, _ = prepared
        gaps = {(g.sheet, g.column, g.row) for g in report.gaps}
        assert ("classes", "class_URI", 2) in gaps
        assert ("classes", "SHACL_independent_class", 2) not in gaps
        assert ("classes", "SHACL_independent_class", 2) in {(g.sheet, g.column, g.row) for g in report.optional_gaps}
        assert ("classes", "SeMPyRO_annotations_IRI", 6) in gaps
        assert ("Dataset", "SeMPyRO_range", 3) in gaps
        assert ("prefixes", "namespace", 2) in gaps
        assert "Required cells to fill in" in report.format()
