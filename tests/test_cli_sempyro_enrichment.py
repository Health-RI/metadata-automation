import openpyxl
import pandas as pd
import pytest
import yaml
from click.testing import CliRunner

from metadata_automation.cli import sempyro_enrichment_apply, sempyro_enrichment_plan

OLD_CLASSES = pd.DataFrame(
    {
        "sheet_name": ["Dataset", "Agent"],
        "class_URI": ["hri:Dataset", "hri:Agent"],
        "SeMPyRO_inherits_from": ["dcat:Dataset", "foaf:Agent"],
        "SeMPyRO_annotations_ontology": ["https://www.w3.org/TR/vocab-dcat-3/", "http://xmlns.com/foaf/spec/"],
        "SeMPyRO_annotations_IRI": ["DCAT.Dataset", "FOAF.Agent"],
    }
)
NEW_CLASSES = pd.DataFrame(
    {
        "sheet_name": ["Dataset", "Agent", "Activity"],
        "class_URI": ["hrishapes:Dataset", "hrishapes:Agent", "hrishapes:Activity"],
        "SeMPyRO_inherits_from": [None, None, None],
        "SeMPyRO_import_classes": [None, None, None],
        "SeMPyRO_add_rdf_model": [None, None, None],
        "SeMPyRO_annotations_ontology": [None, None, None],
        "SeMPyRO_annotations_IRI": [None, None, None],
    }
)
ENRICHMENT = ["SeMPyRO_rdf_term", "SeMPyRO_rdf_type", "SeMPyRO_range"]


def _write(path, sheets):
    with pd.ExcelWriter(path) as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    return path


@pytest.fixture
def old_excel(tmp_path):
    sheets = {
        "classes": OLD_CLASSES,
        "Dataset": pd.DataFrame(
            {
                "Property URI": [
                    "dct:title",
                    "dct:publisher",
                    "dcat:landingPage",
                    "dct:issued",
                    "dcat:theme",
                    "dct:type",
                ],
                "Range": ["rdfs:Literal", "foaf:Agent", "foaf:Document (IRI)", "xsd:dateTime"]
                + ["skos:Concept (IRI)"] * 2,
                "SeMPyRO_rdf_term": ["DCTERMS.title", "DCTERMS.publisher", "DCAT.landingPage", "DCTERMS.issued"]
                + ["DCAT.theme", "DCTERMS.type"],
                "SeMPyRO_rdf_type": ["rdfs_literal", "uri", "uri", "datetime_literal", "uri", "uri"],
                "SeMPyRO_range": ["str, LiteralField", "AnyHttpUrl, HRIAgent", "AnyHttpUrl", "AwareDatetime"]
                + ["AnyHttpUrl", "DatasetTheme"],
            }
        ),
        "Agent": pd.DataFrame(
            {
                "Property URI": ["foaf:name"],
                "Range": ["rdfs:Literal"],
                "SeMPyRO_rdf_term": ["FOAF.name"],
                "SeMPyRO_rdf_type": ["rdfs_literal"],
                "SeMPyRO_range": ["str, LiteralField"],
            }
        ),
    }
    return _write(tmp_path / "old.xlsx", sheets)


@pytest.fixture
def new_excel(tmp_path):
    empty = {c: [None] * 5 for c in ENRICHMENT}
    sheets = {
        "classes": NEW_CLASSES,
        "Dataset": pd.DataFrame(
            {
                "Property URI": ["dct:title", "dct:publisher", "dcat:landingPage", "dct:alternative", "csvw:x"],
                "Range": ["rdfs:Literal", "foaf:Agent (IRI)", "foaf:Document (IRI)", "rdfs:Literal", "csvw:Table"],
                **empty,
            }
        ),
        "Agent": pd.DataFrame(
            {"Property URI": ["foaf:name"], "Range": ["rdfs:Literal"], **{c: [None] for c in ENRICHMENT}}
        ),
        "Activity": pd.DataFrame(
            {"Property URI": ["dct:type"], "Range": ["skos:Concept (IRI)"], **{c: [None] for c in ENRICHMENT}}
        ),
    }
    return _write(tmp_path / "new.xlsx", sheets)


@pytest.fixture
def yamls(tmp_path):
    validation = tmp_path / "validation.yaml"
    validation.write_text(yaml.safe_dump({"classes": {"HRIDataset": {}}}))
    types = tmp_path / "types.yaml"
    types.write_text(yaml.safe_dump({"types": {"AnyHttpUrl": {}, "str": {}}}))
    imports = tmp_path / "imports.yaml"
    imports.write_text(yaml.safe_dump({"hri-Dataset": "import x"}))
    return ["--validation-yaml", validation, "--types-yaml", types, "--imports-yaml", imports]


def _plan_rows(path):
    ws = openpyxl.load_workbook(path)["plan"]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


def _find(rows, sheet, key, column):
    return next(
        r for r in rows if r["sheet"] == sheet and r["key (Property URI / class)"] == key and r["column"] == column
    )


@pytest.fixture
def plan(tmp_path, old_excel, new_excel, yamls):
    path = tmp_path / "plan.xlsx"
    result = CliRunner().invoke(
        sempyro_enrichment_plan,
        ["--previous-excel", str(old_excel), "-i", str(new_excel), "-o", str(path), *map(str, yamls)],
    )
    assert result.exit_code == 0, result.output
    return path


def test_plan_carries_over_unchanged_properties(plan):
    row = _find(_plan_rows(plan), "Dataset", "dct:title", "SeMPyRO_range")
    assert row["status"] == "carried_over"
    assert row["proposed value"] == "str, LiteralField"


def test_plan_flags_changed_range_for_review(plan):
    row = _find(_plan_rows(plan), "Dataset", "dct:publisher", "SeMPyRO_range")
    assert row["status"] == "carried_over_REVIEW"
    assert "foaf:Agent (IRI)" in row["note"]


def test_plan_derives_new_property_values_by_rule(plan):
    rows = _plan_rows(plan)
    term = _find(rows, "Dataset", "dct:alternative", "SeMPyRO_rdf_term")
    assert (term["status"], term["proposed value"]) == ("proposed_by_rule", "DCTERMS.alternative")
    kind = _find(rows, "Dataset", "dct:alternative", "SeMPyRO_rdf_type")
    assert (kind["status"], kind["proposed value"]) == ("proposed_by_rule", "rdfs_literal")


def test_plan_needs_human_for_unknown_prefix_and_range(plan):
    rows = _plan_rows(plan)
    assert _find(rows, "Dataset", "csvw:x", "SeMPyRO_rdf_term")["status"] == "NEEDS_HUMAN"
    assert _find(rows, "Dataset", "csvw:x", "SeMPyRO_range")["status"] == "NEEDS_HUMAN"


def test_plan_proposes_new_class_from_hint(plan):
    row = _find(_plan_rows(plan), "classes", "Activity", "SeMPyRO_annotations_IRI")
    assert (row["status"], row["proposed value"]) == ("proposed_by_rule", "PROV.Activity")
    assert _find(_plan_rows(plan), "classes", "Dataset", "SeMPyRO_annotations_IRI")["proposed value"] == "DCAT.Dataset"


def test_plan_reports_missing_yaml_entries(plan):
    ws = openpyxl.load_workbook(plan)["yaml_coverage"]
    missing = {(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[2] == "missing"}
    assert ("validation_logic.yaml", "HRIActivity") in missing
    assert ("imports.yaml", "hri-Agent") in missing
    assert ("sempyro_types.yaml", "LiteralField") in missing
    assert ("validation_logic.yaml", "HRIDataset") not in missing


def _edit_plan(path, edits):
    """edits: {(sheet, key, column): (final_value, reviewed)}"""
    wb = openpyxl.load_workbook(path)
    ws = wb["plan"]
    header = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        key = (
            row[header["sheet"] - 1].value,
            row[header["key (Property URI / class)"] - 1].value,
            row[header["column"] - 1].value,
        )
        if key in edits:
            final, reviewed = edits[key]
            row[header["final_value (override; blank = accept proposed)"] - 1].value = final
            row[header["reviewed (x)"] - 1].value = reviewed
    wb.save(path)


def _cell(path, sheet, uri, column):
    ws = openpyxl.load_workbook(path)[sheet]
    header = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        if row[header["Property URI"] - 1].value == uri:
            return row[header[column] - 1].value
    raise KeyError(uri)


def test_apply_writes_proposals_and_overrides(tmp_path, new_excel, plan):
    _edit_plan(plan, {("Dataset", "dct:alternative", "SeMPyRO_range"): ("LiteralField", "x")})
    out = tmp_path / "enriched.xlsx"
    result = CliRunner().invoke(sempyro_enrichment_apply, ["-i", str(new_excel), "-p", str(plan), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert _cell(out, "Dataset", "dct:title", "SeMPyRO_range") == "str, LiteralField"
    assert _cell(out, "Dataset", "dct:alternative", "SeMPyRO_range") == "LiteralField"
    classes = openpyxl.load_workbook(out)["classes"]
    header = {c.value: c.column for c in classes[1]}
    assert classes.cell(2, header["SeMPyRO_annotations_IRI"]).value == "DCAT.Dataset"
    # the input stays untouched
    assert _cell(new_excel, "Dataset", "dct:title", "SeMPyRO_range") is None


def test_apply_skips_unreviewed_needs_human_guess(tmp_path, new_excel, plan):
    out = tmp_path / "enriched.xlsx"
    result = CliRunner().invoke(sempyro_enrichment_apply, ["-i", str(new_excel), "-p", str(plan), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert _cell(out, "Activity", "dct:type", "SeMPyRO_range") is None
    assert "Activity / dct:type / SeMPyRO_range" in result.output

    _edit_plan(plan, {("Activity", "dct:type", "SeMPyRO_range"): (None, "x")})
    result = CliRunner().invoke(
        sempyro_enrichment_apply, ["-i", str(new_excel), "-p", str(plan), "-o", str(out), "--overwrite"]
    )
    assert result.exit_code == 0, result.output
    assert _cell(out, "Activity", "dct:type", "SeMPyRO_range") == "AnyHttpUrl"


def test_apply_refuses_to_overwrite_without_flag(tmp_path, new_excel, plan):
    out = tmp_path / "enriched.xlsx"
    out.write_text("existing")
    result = CliRunner().invoke(sempyro_enrichment_apply, ["-i", str(new_excel), "-p", str(plan), "-o", str(out)])
    assert result.exit_code == 1
    assert "--overwrite" in result.output
    result = CliRunner().invoke(sempyro_enrichment_apply, ["-i", str(new_excel), "-p", str(plan), "-o", str(new_excel)])
    assert result.exit_code == 1
    assert "must differ" in result.output
