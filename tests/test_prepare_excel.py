"""Tests for the prepare-excel command and the Excel preparation logic."""

import hashlib

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from metadata_automation.cli import prepare_excel, sempyro, shaclplay
from metadata_automation.excel_prep.prefix_check import check_prefixes
from metadata_automation.excel_prep.preparer import (
    CLASS_SHEET_COLUMNS,
    CLASSES_COLUMNS,
    prepare_workbook,
    relax_cardinality,
)
from metadata_automation.excel_prep.report import Cell, compress_rows
from metadata_automation.excel_prep.workbook import class_sheet_names, data_rows, headers, property_rows, text

# Answers a modeler gives after the first pass: target class and whether it is independent.
CLASS_ANSWERS = {
    "Resource": ("dcat:Resource", False),
    "Catalog": ("dcat:Catalog", True),
    "Dataset": ("dcat:Dataset", True),
    "Agent": ("foaf:Agent", False),
    "Distribution": ("dcat:Distribution", True),
}
NAMESPACES = {
    "hri": "https://w3id.org/health-ri/metadata-vocabulary#",
    "dct": "http://purl.org/dc/terms/",
    "dcat": "http://www.w3.org/ns/dcat#",
    "foaf": "http://xmlns.com/foaf/0.1/",
    "rdfs": "http://www.w3.org/2000/01/rdf-schema#",
    "xsd": "http://www.w3.org/2001/XMLSchema#",
    "skos": "http://www.w3.org/2004/02/skos/core#",
    "dash": "http://datashapes.org/dash#",
}
RDF_NAMESPACES = {"dct": "DCTERMS", "dcat": "DCAT", "foaf": "FOAF"}


@pytest.fixture
def plain_small(test_input_dir):
    return test_input_dir / "plain_small.xlsx"


@pytest.fixture
def plain_metadata(test_input_dir):
    return test_input_dir / "plain_metadata.xlsx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def find_row(ws, column, value):
    cols = headers(ws)
    return next(r for r in range(2, ws.max_row + 1) if text(ws, r, cols[column]) == value)


def cell(ws, row, column):
    return text(ws, row, headers(ws)[column])


def set_cell(ws, row, column, value):
    ws.cell(row=row, column=headers(ws)[column], value=value)


def answer_classes(path, answers=CLASS_ANSWERS):
    """Fill in SHACL_target_ontology_name and SHACL_independent_class, as the modeler would."""
    wb = load_workbook(path)
    ws = wb["classes"]
    for name, (target, independent) in answers.items():
        row = find_row(ws, "sheet_name", name)
        set_cell(ws, row, "SHACL_target_ontology_name", target)
        set_cell(ws, row, "SHACL_independent_class", independent)
    wb.save(path)


@pytest.fixture
def second_pass(plain_small, tmp_path):
    """Prepare, answer the class questions, and prepare again."""
    first = tmp_path / "first.xlsx"
    prepare_workbook(plain_small, first)
    answer_classes(first)
    second = tmp_path / "second.xlsx"
    report = prepare_workbook(first, second)
    return report, load_workbook(second), second


class TestFirstPass:
    @pytest.fixture
    def prepared(self, plain_small, tmp_path):
        out = tmp_path / "prepared.xlsx"
        report = prepare_workbook(plain_small, out)
        return report, load_workbook(out)

    def test_source_is_untouched(self, plain_small, tmp_path):
        before = sha(plain_small)
        prepare_workbook(plain_small, tmp_path / "out.xlsx")
        assert sha(plain_small) == before

    def test_refuses_to_overwrite_source(self, plain_small):
        with pytest.raises(ValueError, match="differ"):
            prepare_workbook(plain_small, plain_small)

    def test_general_sheets_added_after_info(self, prepared):
        _, wb = prepared
        assert wb.sheetnames[:4] == ["Info", "prefixes", "classes", "Resource"]
        assert wb["Info"]["A2"].value == "Plain test model"

    def test_classes_sheet_has_row_per_class_sheet(self, prepared):
        _, wb = prepared
        ws = wb["classes"]
        cols = headers(ws)
        assert list(cols) == CLASSES_COLUMNS
        names = [text(ws, r, cols["sheet_name"]) for r in data_rows(ws, cols, ("sheet_name",))]
        assert names == class_sheet_names(wb) == ["Resource", "Catalog", "Dataset", "Agent", "Distribution"]

    def test_class_sheets_have_pipeline_columns(self, prepared):
        _, wb = prepared
        for name in class_sheet_names(wb):
            assert set(CLASS_SHEET_COLUMNS) <= set(headers(wb[name]))

    def test_added_columns_copy_first_column_format(self, plain_small, tmp_path):
        styled = tmp_path / "styled.xlsx"
        wb = load_workbook(plain_small)
        ws = wb["Dataset"]
        ws["A1"].fill = PatternFill("solid", fgColor="FFC000")
        ws["A1"].font = Font(bold=True, color="FF0000")
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 31
        wb.save(styled)

        out = tmp_path / "prepared.xlsx"
        prepare_workbook(styled, out)
        ws = load_workbook(out)["Dataset"]
        for name in CLASS_SHEET_COLUMNS:
            col = headers(ws)[name]
            header, data = ws.cell(row=1, column=col), ws.cell(row=2, column=col)
            assert header.fill.fgColor.rgb == "00FFC000"
            assert header.font.b and header.font.color.rgb == "00FF0000"
            assert data.alignment.wrap_text
            assert ws.column_dimensions[get_column_letter(col)].width == 31

    def test_prefixes_sheet_lists_used_prefixes_without_namespace(self, prepared):
        report, wb = prepared
        ws = wb["prefixes"]
        cols = headers(ws)
        prefixes = {text(ws, r, cols["prefix"]) for r in data_rows(ws, cols, ("prefix",))}
        assert prefixes == {"dcat", "dct", "foaf", "rdfs", "skos", "xsd"}
        assert all(not text(ws, r, cols["namespace"]) for r in data_rows(ws, cols, ("prefix",)))
        assert any("'dct' has no registered namespace" in f.message for f in report.prefix_issues)

    def test_nothing_is_guessed(self, prepared):
        report, wb = prepared
        ws = wb["classes"]
        for row in data_rows(ws, headers(ws), ("sheet_name",)):
            for column in ("class_URI", "SHACL_target_ontology_name", "SHACL_independent_class"):
                assert not cell(ws, row, column)
        assert cell(wb["Catalog"], 3, "Range") == "dcat:Dataset"
        assert any("range rules skipped" in f.message for f in report.issues)

    def test_report_lists_gaps(self, prepared):
        report, _ = prepared
        gaps = {(g.sheet, g.column, g.row) for g in report.gaps}
        assert ("classes", "SHACL_independent_class", 2) in gaps
        assert ("classes", "SeMPyRO_annotations_IRI", 6) in gaps
        assert ("Dataset", "SeMPyRO_range", 3) in gaps
        assert ("prefixes", "namespace", 2) in gaps
        assert "Required cells to fill in" in report.format()


class TestSecondPass:
    def test_class_uri_derived_from_target(self, second_pass):
        _, wb, _ = second_pass
        ws = wb["classes"]
        assert cell(ws, find_row(ws, "sheet_name", "Catalog"), "class_URI") == "hri:Catalog"

    def test_independent_class_range_gets_iri_and_relaxed_cardinality(self, second_pass):
        _, wb, _ = second_pass
        catalog, dataset = wb["Catalog"], wb["Dataset"]
        row = find_row(catalog, "Property URI", "dcat:dataset")
        assert cell(catalog, row, "Range") == "dcat:Dataset (IRI)"
        assert cell(catalog, row, "Cardinality") == "0..n"
        row = find_row(dataset, "Property URI", "dcat:distribution")
        assert cell(dataset, row, "Range") == "dcat:Distribution (IRI)"
        assert cell(dataset, row, "Cardinality") == "0..n"
        assert not cell(dataset, row, "SHACL_sh:node")

    def test_dependent_class_range_gets_node_shape(self, second_pass):
        _, wb, _ = second_pass
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        assert cell(ws, row, "Range") == "foaf:Agent"
        assert cell(ws, row, "Cardinality") == "1"
        assert cell(ws, row, "SHACL_sh:node") == "hri:AgentShape"

    @pytest.mark.parametrize(
        "sheet, uri, expected",
        [
            ("Catalog", "dct:license", "dct:LicenseDocument (IRI)"),
            ("Agent", "dct:type", "skos:Concept (IRI)"),
            ("Dataset", "dcat:theme", "skos:Concept (IRI)"),
            ("Catalog", "dct:modified", "xsd:dateTime"),
            ("Catalog", "dct:title", "rdfs:Literal"),
            ("Distribution", "dcat:accessURL", "rdfs:Resource (IRI)"),
        ],
    )
    def test_out_of_model_ranges(self, second_pass, sheet, uri, expected):
        _, wb, _ = second_pass
        ws = wb[sheet]
        assert cell(ws, find_row(ws, "Property URI", uri), "Range") == expected

    @pytest.mark.parametrize(
        "uri, viewer, editor",
        [
            ("dcat:dataset", "dash:URIViewer", "dash:URIEditor"),
            ("dct:publisher", "dash:DetailsViewer", "dash:BlankNodeEditor"),
            ("dct:modified", "dash:LiteralViewer", "dash:DateTimePickerEditor"),
            ("dct:description", "dash:LiteralViewer", "dash:TextAreaEditor"),
            ("dct:title", "dash:LiteralViewer", "dash:TextFieldEditor"),
        ],
    )
    def test_empty_dash_cells_are_filled(self, second_pass, uri, viewer, editor):
        _, wb, _ = second_pass
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", uri)
        assert cell(ws, row, "SHACL_dash:viewer") == viewer
        assert cell(ws, row, "SHACL_dash:editor") == editor

    def test_is_idempotent(self, second_pass, tmp_path):
        _, _, path = second_pass
        report = prepare_workbook(path, tmp_path / "third.xlsx")
        assert report.changes == []

    def test_no_title_issue_when_independent_classes_have_one(self, second_pass):
        report, _, _ = second_pass
        assert not any("dct:title" in f.message for f in report.issues)


class TestInit:
    @pytest.fixture
    def answered(self, plain_small, tmp_path):
        """A first-pass output with the class questions answered, so every rule could apply."""
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first)
        return first

    def test_adds_structure_and_prefixes(self, plain_small, tmp_path):
        out = tmp_path / "init.xlsx"
        report = prepare_workbook(plain_small, out, init=True)
        wb = load_workbook(out)
        assert wb.sheetnames[:4] == ["Info", "prefixes", "classes", "Resource"]
        for name in class_sheet_names(wb):
            assert set(CLASS_SHEET_COLUMNS) <= set(headers(wb[name]))
        assert any(c.cell.sheet == "prefixes" for c in report.changes)
        assert report.prefix_issues

    def test_leaves_cell_values_untouched(self, answered, tmp_path):
        out = tmp_path / "init.xlsx"
        report = prepare_workbook(answered, out, init=True)
        before, after = load_workbook(answered), load_workbook(out)
        for name in before.sheetnames:
            assert [[c.value for c in r] for r in before[name].iter_rows()] == [
                [c.value for c in r] for r in after[name].iter_rows()
            ]
        assert report.changes == []
        assert not any("range rules skipped" in f.message for f in report.issues)

    def test_still_runs_checks(self, answered, tmp_path):
        wb = load_workbook(answered)
        ws = wb["Catalog"]
        ws.delete_rows(find_row(ws, "Property URI", "dct:title"))
        wb.save(answered)
        report = prepare_workbook(answered, tmp_path / "init.xlsx", init=True)
        assert any(f.cell.sheet == "Catalog" and "dct:title" in f.message for f in report.issues)


class TestIssues:
    def test_dependent_range_with_iri_is_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Catalog"]
        set_cell(ws, find_row(ws, "Property URI", "dct:publisher"), "Range", "foaf:Agent (IRI)")
        wb.save(first)

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        assert any(f.cell.sheet == "Catalog" and "dependent class" in f.message for f in report.issues)

    def test_differing_dash_value_is_reported_not_overwritten(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first)
        wb = load_workbook(first)
        ws = wb["Catalog"]
        row = find_row(ws, "Property URI", "dct:publisher")
        set_cell(ws, row, "SHACL_dash:viewer", "dash:URIViewer")
        wb.save(first)

        second = tmp_path / "second.xlsx"
        report = prepare_workbook(first, second)
        assert any(
            f.cell == Cell("Catalog", row, "SHACL_dash:viewer")
            and f.message == "expected 'dash:DetailsViewer', found 'dash:URIViewer'"
            for f in report.issues
        )
        ws = load_workbook(second)["Catalog"]
        assert cell(ws, row, "SHACL_dash:viewer") == "dash:URIViewer"
        assert cell(ws, row, "SHACL_dash:editor") == "dash:BlankNodeEditor"

    def test_no_dash_values_before_targets_are_known(self, plain_small, tmp_path):
        out = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, out)
        ws = load_workbook(out)["Catalog"]
        assert not cell(ws, find_row(ws, "Property URI", "dct:title"), "SHACL_dash:viewer")

    def test_independent_class_without_title_is_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        answer_classes(first, {**CLASS_ANSWERS, "Agent": ("foaf:Agent", True)})

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        assert any(f.cell.sheet == "Agent" and "no dct:title" in f.message for f in report.issues)

    def test_invalid_flag_and_unknown_sheet_are_reported(self, plain_small, tmp_path):
        first = tmp_path / "first.xlsx"
        prepare_workbook(plain_small, first)
        wb = load_workbook(first)
        ws = wb["classes"]
        set_cell(ws, find_row(ws, "sheet_name", "Agent"), "SHACL_independent_class", "maybe")
        ws.cell(row=ws.max_row + 1, column=1, value="Ghost")
        wb.save(first)

        report = prepare_workbook(first, tmp_path / "second.xlsx")
        messages = [f.message for f in report.issues]
        assert "expected TRUE or FALSE, got 'maybe'" in messages
        assert "sheet 'Ghost' does not exist" in messages


class TestPrefixCheck:
    @pytest.fixture
    def findings(self):
        wb = Workbook()
        prefixes = wb.active
        prefixes.title = "prefixes"
        for row in [
            ("prefix", "namespace"),
            ("dct", "http://purl.org/dc/terms/"),
            ("dcterms", "http://purl.org/dc/terms/"),
            ("xsd", "http://www.w3.org/2001/XMLSchema#"),
            ("xsd", "http://example.org/xsd#"),
            ("DCAT", "http://www.w3.org/ns/dcat#"),
        ]:
            prefixes.append(row)
        sheet = wb.create_sheet("Thing")
        for row in [
            ("Property label", "Property URI", "Range"),
            ("legislation", "dcat:legislation", "eli:LegalResource (IRI)"),
            ("page", "http://purl.org/dc/terms/page ", "rdfs:Resource  (IRI)"),
            ("name", "dct:name", "Literal"),
        ]:
            sheet.append(row)
        return [f"{f.cell} | {f.message}" for f in check_prefixes(wb)]

    @pytest.mark.parametrize(
        "expected",
        [
            "'xsd' registered in rows 4, 5 with conflicting namespaces",
            "namespace http://purl.org/dc/terms/ registered under several prefixes: dct, dcterms",
            "prefix spelled in different cases: DCAT, dcat",
            "'eli' has no registered namespace, used in: Thing / Range / row 2",
            "'dcat' has no registered namespace",
            "Thing / Range / row 3 | stray whitespace in 'rdfs:Resource  (IRI)'",
            "Thing / Property URI / row 3 | stray whitespace",
            "full IRI 'http://purl.org/dc/terms/page', use CURIE 'dct:page'",
            "Thing / Range / row 4 | no prefix in 'Literal'",
        ],
    )
    def test_finds(self, findings, expected):
        assert any(expected in f for f in findings), findings


@pytest.mark.parametrize(
    "given, expected",
    [("1", "0..1"), ("1..1", "0..1"), ("1..n", "0..n"), ("2..5", "0..5"), ("0..n", "0..n"), ("0..1", "0..1")],
)
def test_relax_cardinality(given, expected):
    assert relax_cardinality(given) == expected


def test_relax_cardinality_unparsable():
    assert relax_cardinality("one or more") is None


def test_compress_rows():
    assert compress_rows([7, 2, 3, 4, 9, 10]) == "2-4, 7, 9-10"


def test_plain_metadata_gets_all_classes(plain_metadata, tmp_path):
    out = tmp_path / "prepared.xlsx"
    prepare_workbook(plain_metadata, out)
    wb = load_workbook(out)
    ws = wb["classes"]
    cols = headers(ws)
    names = [text(ws, r, cols["sheet_name"]) for r in data_rows(ws, cols, ("sheet_name",))]
    assert names == class_sheet_names(wb)
    assert "Resource" in names and "CatalogueRecord" in names


def test_cli_prints_report(runner, plain_small, tmp_path):
    out = tmp_path / "prepared.xlsx"
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert "Required cells to fill in" in result.output
    assert out.exists()


def test_cli_init_flag(runner, plain_small, tmp_path):
    out = tmp_path / "prepared.xlsx"
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(out), "--init"])
    assert result.exit_code == 0, result.output
    assert "Prefix check" in result.output
    assert "derived from range" not in result.output


def test_cli_rejects_same_output(runner, plain_small):
    result = runner.invoke(prepare_excel, ["-i", str(plain_small), "-o", str(plain_small)])
    assert result.exit_code == 1
    assert "must differ" in result.output


def fill_reported_gaps(path):
    """Fill every required gap with a valid value, as the modeler would after reading the report."""
    wb = load_workbook(path)
    ws = wb["prefixes"]
    for row in data_rows(ws, headers(ws), ("prefix",)):
        set_cell(ws, row, "namespace", NAMESPACES[cell(ws, row, "prefix")])

    ws = wb["classes"]
    for row in data_rows(ws, headers(ws), ("sheet_name",)):
        set_cell(ws, row, "SeMPyRO_annotations_ontology", "https://www.w3.org/TR/vocab-dcat-3/")
        set_cell(ws, row, "SeMPyRO_annotations_IRI", f"https://example.org/{cell(ws, row, 'sheet_name')}")
        set_cell(ws, row, "SeMPyRO_add_rdf_model", True)

    for name in class_sheet_names(wb):
        ws = wb[name]
        for row in property_rows(ws, headers(ws)):
            prefix, local = cell(ws, row, "Property URI").split(":")
            literal = cell(ws, row, "Range").startswith(("rdfs:Literal", "xsd:"))
            set_cell(ws, row, "SeMPyRO_rdf_term", f"{RDF_NAMESPACES[prefix]}.{local}")
            set_cell(ws, row, "SeMPyRO_rdf_type", "rdfs_literal" if literal else "uri")
            set_cell(ws, row, "SeMPyRO_range", "LiteralField, str" if literal else "AnyHttpUrl")
    wb.save(path)


def test_end_to_end_shaclplay_and_sempyro(runner, second_pass, tmp_path):
    _, _, path = second_pass
    fill_reported_gaps(path)
    final = tmp_path / "final.xlsx"
    report = prepare_workbook(path, final)
    assert report.gaps == []

    result = runner.invoke(shaclplay, ["-i", str(final), "-o", str(tmp_path / "shaclplay")])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "shaclplay" / "SHACL-catalog.xlsx").exists()

    imports = tmp_path / "imports.yaml"
    header = (
        "    from pydantic import AnyHttpUrl, ConfigDict, Field\n"
        "    from rdflib.namespace import DCAT, DCTERMS, FOAF\n"
        "    from sempyro import LiteralField, RDFModel\n"
    )
    imports.write_text("".join(f"hri-{name}: |\n{header}\n" for name in CLASS_ANSWERS))
    result = runner.invoke(
        sempyro,
        [
            "-i",
            str(final),
            "--linkml-output-path",
            str(tmp_path / "linkml"),
            "--sempyro-output-path",
            str(tmp_path / "sempyro"),
            "--imports-path",
            str(imports),
        ],
    )
    assert result.exit_code == 0, result.output
    assert "Successfully generated: 5 classes" in result.output
