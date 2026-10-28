"""Carry the SeMPyRO enrichment of one model Excel over to a newer one, through a plan a person reviews.

``build_plan`` compares the enriched Excel of the previous model version with the new, plain Excel and lists every
SeMPyRO cell to fill, with a proposed value, a status and the reason. The plan is an Excel file a reviewer edits.
``apply_plan`` then writes the reviewed values into a copy of the new Excel.

The SeMPyRO columns of the v2.0.2 model were filled in by hand, so only part of them follows a rule. The plan keeps
that visible: values copied for an unchanged property are ``carried_over``, values derived from patterns in the
previous version are ``proposed_by_rule`` and the rest is ``NEEDS_HUMAN``.
"""

import collections
import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from metadata_automation.excel_prep.drawings import restore_drawings

SKIP_SHEETS = {"info", "user guide", "prefixes", "classes"}
CLASS_COLUMNS = [
    "SeMPyRO_inherits_from",
    "SeMPyRO_import_classes",
    "SeMPyRO_annotations_ontology",
    "SeMPyRO_annotations_IRI",
    "SeMPyRO_add_rdf_model",
]
PROPERTY_COLUMNS = ["SeMPyRO_rdf_term", "SeMPyRO_rdf_type", "SeMPyRO_range"]

CARRIED = "carried_over"  # same property and Range as in the previous version
CARRIED_REVIEW = "carried_over_REVIEW"  # same property, but the Range or URI changed
RULE = "proposed_by_rule"  # derived from patterns in the previous version
NEEDS_HUMAN = "NEEDS_HUMAN"  # no reliable rule; a reviewer has to decide
STATUS_ORDER = {NEEDS_HUMAN: 0, CARRIED_REVIEW: 1, RULE: 2, CARRIED: 3}
STATUS_COLOR = {CARRIED: "E2F0D9", CARRIED_REVIEW: "FFF2CC", RULE: "DDEBF7", NEEDS_HUMAN: "F8CBAD"}
# A value is proposed by rule when this share of the previous rows with the same Range agree on it.
RULE_AGREEMENT = 0.9

PLAN_SHEET = "plan"
FINAL_COLUMN = "final_value (override; blank = accept proposed)"
REVIEWED_COLUMN = "reviewed (x)"
PLAN_COLUMNS = [
    "id",
    "level",
    "sheet",
    "key (Property URI / class)",
    "Property label",
    "Range (v2.1.0)",
    "column",
    "row in sheet",
    "current value",
    "proposed value",
    "status",
    "source / why",
    "note",
    "alternatives seen in previous version",
    FINAL_COLUMN,
    REVIEWED_COLUMN,
]
PLAN_WIDTHS = [6, 9, 18, 34, 28, 26, 28, 8, 20, 36, 20, 50, 40, 36, 36, 10]

# Proposals for classes that are new in a model version; upstream vocabulary names, to be verified by the reviewer.
NEW_CLASS_HINTS: dict[str, dict[str, Any]] = {
    "Activity": {
        "SeMPyRO_add_rdf_model": True,
        "SeMPyRO_annotations_ontology": "https://www.w3.org/TR/prov-o/",
        "SeMPyRO_annotations_IRI": "PROV.Activity",
        "note": "SeMPyRO already ships prov.Activity",
    },
    "ProvenanceStatement": {
        "SeMPyRO_add_rdf_model": True,
        "SeMPyRO_annotations_ontology": "https://www.dublincore.org/specifications/dublin-core/dcmi-terms/",
        "SeMPyRO_annotations_IRI": "DCTERMS.ProvenanceStatement",
        "note": "Not shipped by SeMPyRO; check the IRI and the ontology URL",
    },
    **{
        name: {
            "SeMPyRO_add_rdf_model": True,
            "SeMPyRO_annotations_ontology": "https://www.w3.org/TR/tabular-metadata/",
            "SeMPyRO_annotations_IRI": f"CSVW.{name}",
            "note": "SeMPyRO has no CSVW namespace or classes; add them before generating this class",
        }
        for name in ("TableGroup", "Table", "Column")
    },
}


@dataclass
class PlanItem:
    """One SeMPyRO cell to fill."""

    level: str  # "class" or "property"
    sheet: str
    key: str  # class name or Property URI
    label: str
    row: int
    column: str
    proposed: Any
    status: str
    source: str
    range_: str = ""
    current: Any = None
    note: str = ""
    alternatives: str = ""
    class_uri: str = ""
    id: int = 0


@dataclass
class Patterns:
    """What the previous version's enrichment does, per prefix and per Range."""

    namespace: dict[str, collections.Counter] = field(
        default_factory=lambda: collections.defaultdict(collections.Counter)
    )
    by_range: dict[str, dict[str, collections.Counter]] = field(
        default_factory=lambda: {c: collections.defaultdict(collections.Counter) for c in PROPERTY_COLUMNS[1:]}
    )


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value not in (None, "") else ""


def _clean(value: Any) -> Any:
    return value.strip() if isinstance(value, str) else value


def _read_table(ws: Worksheet) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    header = [_norm(h) for h in rows[0]]
    table = []
    for number, row in enumerate(rows[1:], start=2):
        if any(cell not in (None, "") for cell in row):
            table.append({**{h: v for h, v in zip(header, row, strict=False) if h}, "_row": number})
    return table


def _class_sheets(wb: openpyxl.Workbook) -> list[Worksheet]:
    return [ws for ws in wb.worksheets if ws.title.lower() not in SKIP_SHEETS]


def _learn_patterns(old_wb: openpyxl.Workbook) -> Patterns:
    patterns = Patterns()
    for ws in _class_sheets(old_wb):
        for row in _read_table(ws):
            uri = _norm(row.get("Property URI"))
            if ":" not in uri:
                continue
            term = _norm(row.get("SeMPyRO_rdf_term"))
            if term:
                patterns.namespace[uri.split(":")[0]][term.split(".")[0]] += 1
            range_ = _norm(row.get("Range"))
            for column in PROPERTY_COLUMNS[1:]:
                value = _norm(row.get(column))
                if value:
                    patterns.by_range[column][range_][value] += 1
    return patterns


def _propose_new(column: str, uri: str, range_: str, patterns: Patterns) -> tuple[Any, str, str, str]:
    """Value, status, source and alternatives for a property that did not exist in the previous version."""
    if column == "SeMPyRO_rdf_term":
        prefix, _, local = uri.partition(":")
        counter = patterns.namespace.get(prefix)
        if not counter:
            return None, NEEDS_HUMAN, f"prefix '{prefix}' was not used before: pick the SeMPyRO/rdflib namespace", ""
        top = counter.most_common(1)[0][0]
        alternatives = ", ".join(f"{k} ({v}x)" for k, v in counter.most_common()) if len(counter) > 1 else ""
        return f"{top}.{local}", RULE, f"namespace of prefix '{prefix}' in the previous version: {top}", alternatives
    counter = patterns.by_range[column].get(range_)
    if not counter:
        return None, NEEDS_HUMAN, f"Range '{range_}' was not used before", ""
    top, count = counter.most_common(1)[0]
    total = sum(counter.values())
    alternatives = ", ".join(f"{k} ({v}x)" for k, v in counter.most_common())
    if count / total >= RULE_AGREEMENT:
        return top, RULE, f"Range '{range_}' gives '{top}' in {count}/{total} rows of the previous version", ""
    return top, NEEDS_HUMAN, f"most common for Range '{range_}' ({count}/{total}); ambiguous", alternatives


def _class_items(old_wb: openpyxl.Workbook, new_wb: openpyxl.Workbook) -> list[PlanItem]:
    old_rows = {_norm(r["sheet_name"]).lower(): r for r in _read_table(old_wb["classes"])}
    items = []
    for row in _read_table(new_wb["classes"]):
        name = _norm(row["sheet_name"])
        old = old_rows.get(name.lower())
        common = {
            "level": "class",
            "sheet": "classes",
            "key": name,
            "label": name,
            "row": row["_row"],
            "class_uri": _norm(row.get("class_URI")),
        }
        hint = NEW_CLASS_HINTS.get(name, {})
        for column in CLASS_COLUMNS:
            current = row.get(column)
            if old is not None:
                items.append(
                    PlanItem(
                        **common,
                        column=column,
                        current=current,
                        proposed=_clean(old.get(column)),
                        status=CARRIED,
                        source=f"previous version, classes row '{name}'",
                    )
                )
                continue
            value = hint.get(column)
            items.append(
                PlanItem(
                    **common,
                    column=column,
                    current=current,
                    proposed=value,
                    status=RULE if value not in (None, "") else NEEDS_HUMAN,
                    source="new class; proposal from the upstream vocabulary" if value else "new class",
                    note=hint.get("note", ""),
                )
            )
        old_local = _norm(old.get("class_URI")).split(":")[-1] if old else ""
        if old_local and old_local != common["class_uri"].split(":")[-1]:
            items.append(
                PlanItem(
                    **common,
                    column="class_URI",
                    current=common["class_uri"],
                    proposed=common["class_uri"],
                    status=CARRIED_REVIEW,
                    source="class_URI local name differs from the previous version",
                    note=f"previous: {_norm(old.get('class_URI'))}",
                )
            )
    return items


def _match_previous(candidates: list[dict[str, Any]], label: str) -> dict[str, Any] | None:
    if len(candidates) > 1:
        candidates = [c for c in candidates if _norm(c.get("Property label")) == label] or candidates
    return candidates[0] if candidates else None


def _carried_item(common: dict[str, Any], column: str, current: Any, old: dict[str, Any]) -> PlanItem:
    """The previous value of ``column``, flagged for review when the Range or the URI spelling changed."""
    notes = []
    old_range = _norm(old.get("Range"))
    if old_range != common["range_"]:
        notes.append(f"Range changed: '{old_range}' -> '{common['range_']}'")
    if _norm(old.get("Property URI")) != common["key"]:
        notes.append(f"URI spelling changed: {_norm(old.get('Property URI'))} -> {common['key']}")
    value = _clean(old.get(column))
    if value in (None, ""):
        status, source = NEEDS_HUMAN, "empty in the previous version"
    else:
        status, source = (CARRIED_REVIEW if notes else CARRIED), f"previous version, sheet '{common['sheet']}'"
    return PlanItem(
        **common, column=column, current=current, proposed=value, status=status, source=source, note="; ".join(notes)
    )


def _property_items(old_wb: openpyxl.Workbook, new_wb: openpyxl.Workbook, patterns: Patterns) -> list[PlanItem]:
    old_sheets = {ws.title.lower(): ws for ws in _class_sheets(old_wb)}
    items = []
    for ws in _class_sheets(new_wb):
        previous = collections.defaultdict(list)
        if ws.title.lower() in old_sheets:
            for row in _read_table(old_sheets[ws.title.lower()]):
                if _norm(row.get("Property URI")):
                    previous[_norm(row["Property URI"]).lower()].append(row)
        note = "" if ws.title.lower() in old_sheets else "new sheet"
        for row in _read_table(ws):
            uri = _norm(row.get("Property URI"))
            if not uri:
                continue
            label = _norm(row.get("Property label"))
            common = {
                "level": "property",
                "sheet": ws.title,
                "key": uri,
                "label": label,
                "row": row["_row"],
                "range_": _norm(row.get("Range")),
            }
            old = _match_previous(previous.get(uri.lower(), []), label)
            for column in PROPERTY_COLUMNS:
                if old is not None:
                    items.append(_carried_item(common, column, row.get(column), old))
                    continue
                value, status, source, alternatives = _propose_new(column, uri, common["range_"], patterns)
                items.append(
                    PlanItem(
                        **common,
                        column=column,
                        current=row.get(column),
                        proposed=value,
                        status=status,
                        source=source,
                        alternatives=alternatives,
                        note=note or "new property",
                    )
                )
    return items


def build_plan(old_excel: Path, new_excel: Path) -> list[PlanItem]:
    """All SeMPyRO cells of ``new_excel`` with a proposed value, taken from or derived from ``old_excel``."""
    old_wb = openpyxl.load_workbook(old_excel, data_only=True)
    new_wb = openpyxl.load_workbook(new_excel, data_only=True)
    items = _class_items(old_wb, new_wb) + _property_items(old_wb, new_wb, _learn_patterns(old_wb))
    for number, item in enumerate(items, start=1):
        item.id = number
    return items


def yaml_coverage(
    items: list[PlanItem], validation_yaml: Path, types_yaml: Path, imports_yaml: Path, namespace: str
) -> list[tuple[str, str, str, str]]:
    """Which Pydantic classes and SeMPyRO types have no entry in the three hand-maintained YAML files."""
    validators = (yaml.safe_load(validation_yaml.read_text()) or {}).get("classes", {})
    types = (yaml.safe_load(types_yaml.read_text()) or {}).get("types", {})
    imports = yaml.safe_load(imports_yaml.read_text()) or {}
    rows = []
    classes = sorted({i.class_uri.split(":")[-1] for i in items if i.level == "class" and i.class_uri})
    for name in classes:
        pydantic = f"{namespace.upper()}{name}"
        rows.append(
            (
                "validation_logic.yaml",
                pydantic,
                "present" if pydantic in validators else "missing",
                "a class without validators may legitimately have no entry",
            )
        )
        key = f"{namespace}-{name}"
        rows.append(("imports.yaml", key, "present" if key in imports else "missing", ""))
    used = collections.Counter(
        t.strip() for i in items if i.column == "SeMPyRO_range" and i.proposed for t in str(i.proposed).split(",")
    )
    for type_name, count in sorted(used.items()):
        if type_name:
            state = "present" if type_name in types else "missing"
            rows.append(("sempyro_types.yaml", type_name, state, f"used by {count} cells"))
    return rows


def _write_readme(wb: openpyxl.Workbook, old_excel: Path, new_excel: Path) -> None:
    readme = wb.active
    readme.title = "README"
    for line in [
        "SeMPyRO enrichment plan",
        "",
        f"Values are carried over from: {old_excel.name}",
        f"Values are written into:      {new_excel.name}",
        "",
        "Review the sheet 'plan' (filter on 'status'):",
        "  carried_over         same property and Range as before. Skim.",
        "  carried_over_REVIEW  same property, but the Range or URI changed. Check the old value still applies.",
        "  proposed_by_rule     derived from patterns in the previous version. Check.",
        "  NEEDS_HUMAN          no reliable rule. Fill in 'final_value'. A shown guess is only used when you",
        "                       fill in 'final_value' or tick 'reviewed (x)'.",
        "'final_value' overrides the proposal; leave it blank to accept the proposal.",
        "",
        "The sheet 'yaml_coverage' lists entries missing in validation_logic.yaml, imports.yaml, sempyro_types.yaml.",
        "Then write the values into the Excel with: metadata-automation sempyro-enrichment-apply",
    ]:
        readme.append([line])
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 110


def _write_plan_sheet(wb: openpyxl.Workbook, items: list[PlanItem]) -> None:
    plan = wb.create_sheet(PLAN_SHEET)
    plan.append(PLAN_COLUMNS)
    for cell in plan[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for item in sorted(items, key=lambda i: (STATUS_ORDER[i.status], i.id)):
        plan.append(
            [item.id, item.level, item.sheet, item.key, item.label, item.range_, item.column, item.row, item.current]
            + [item.proposed, item.status, item.source, item.note, item.alternatives, None, None]
        )
        plan.cell(plan.max_row, 11).fill = PatternFill("solid", fgColor=STATUS_COLOR[item.status])
        plan.cell(plan.max_row, 15).fill = PatternFill("solid", fgColor="FFFFE0")
    plan.auto_filter.ref = plan.dimensions
    plan.freeze_panes = "E2"
    for letter, width in zip("ABCDEFGHIJKLMNOP", PLAN_WIDTHS, strict=True):
        plan.column_dimensions[letter].width = width
    for row in plan.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ticks = DataValidation(type="list", formula1='"x"', allow_blank=True)
    plan.add_data_validation(ticks)
    ticks.add(f"P2:P{plan.max_row}")


def _write_coverage_sheet(wb: openpyxl.Workbook, coverage: list[tuple[str, str, str, str]]) -> None:
    sheet = wb.create_sheet("yaml_coverage")
    sheet.append(["file", "entry", "state", "note"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in coverage:
        sheet.append(list(row))
        if row[2] == "missing":
            sheet.cell(sheet.max_row, 3).fill = PatternFill("solid", fgColor=STATUS_COLOR[NEEDS_HUMAN])
    for letter, width in zip("ABCD", [26, 36, 12, 70], strict=True):
        sheet.column_dimensions[letter].width = width


def write_plan(
    items: list[PlanItem], coverage: list[tuple[str, str, str, str]], output: Path, old_excel: Path, new_excel: Path
) -> collections.Counter:
    """Write the plan as an Excel file a reviewer can edit. Returns the number of items per (level, status)."""
    wb = openpyxl.Workbook()
    _write_readme(wb, old_excel, new_excel)
    _write_plan_sheet(wb, items)
    _write_coverage_sheet(wb, coverage)
    counts = collections.Counter((i.level, i.status) for i in items)
    summary = wb.create_sheet("summary")
    summary.append(["level", "status", "count"])
    for (level, status), count in sorted(counts.items()):
        summary.append([level, status, count])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return counts


def _save_keeping_drawings(wb: openpyxl.Workbook, source: Path, output: Path) -> None:
    """Save through a temporary file, so a failed save leaves ``output`` intact, and restore ``source``'s drawings."""
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=output.suffix, dir=output.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        lost = restore_drawings(source, Path(tmp))
        if lost:
            raise ValueError(f"Drawings of sheets {', '.join(lost)} could not be kept.")
        os.replace(tmp, output)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _value_to_write(row: tuple, index: dict[str, int]) -> Any:
    """``final_value``, else the proposal; an unreviewed ``NEEDS_HUMAN`` proposal is not used."""
    final = row[index[FINAL_COLUMN]]
    if final not in (None, ""):
        return final
    if row[index["status"]] == NEEDS_HUMAN and not row[index[REVIEWED_COLUMN]]:
        return None
    return row[index["proposed value"]]


def apply_plan(new_excel: Path, plan_excel: Path, output: Path, overwrite: bool = False) -> tuple[int, list[tuple]]:
    """Write the reviewed plan into a copy of ``new_excel``, keeping its drawings.

    A value is ``final_value`` when filled in, else the proposal. A ``NEEDS_HUMAN`` proposal is only used when the
    reviewer filled in ``final_value`` or ticked ``reviewed``. Returns the number of cells written and the
    ``NEEDS_HUMAN`` cells that stay empty.
    """
    if output.exists():
        if os.path.samefile(new_excel, output):
            raise ValueError("Output path must differ from the input Excel; the input is left untouched.")
        if not overwrite:
            raise ValueError(f"Output file {output} already exists; pass --overwrite to replace it.")
    rows = list(openpyxl.load_workbook(plan_excel, data_only=True)[PLAN_SHEET].iter_rows(values_only=True))
    index = {name: i for i, name in enumerate(rows[0])}
    wb = openpyxl.load_workbook(new_excel)
    written, empty = 0, []
    for row in rows[1:]:
        column = row[index["column"]]
        if column == "class_URI":
            continue  # reported for review only; never written
        value = _value_to_write(row, index)
        sheet = row[index["sheet"]]
        if value in (None, ""):
            if row[index["status"]] == NEEDS_HUMAN:
                empty.append((sheet, row[index["key (Property URI / class)"]], column))
            continue
        headers = {_norm(c.value): c.column for c in wb[sheet][1] if c.value}
        if column not in headers:
            raise ValueError(f"Column {column} is missing in sheet {sheet}.")
        wb[sheet].cell(row=int(row[index["row in sheet"]]), column=headers[column]).value = _clean(value)
        written += 1
    _save_keeping_drawings(wb, new_excel, output)
    return written, empty
