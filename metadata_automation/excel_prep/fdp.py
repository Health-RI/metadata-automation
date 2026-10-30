"""Derive the FDP variant of a metadata Excel: independent classes are referenced by IRI.

The FDP lists a class with ``SHACL_independent_class`` TRUE as a resource of its own. A property pointing at such a
class therefore refers to it by IRI instead of nesting its properties. The result is a new workbook; the source is
never modified.
"""

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from metadata_automation.excel_prep.class_sheet import set_node
from metadata_automation.excel_prep.columns import FDP_SHEETS
from metadata_automation.excel_prep.dash import URI
from metadata_automation.excel_prep.gaps import add_property_uris
from metadata_automation.excel_prep.preparer import prepare_in_memory, resolve_output, save_workbook
from metadata_automation.excel_prep.report import Report
from metadata_automation.excel_prep.structure import ModelClass
from metadata_automation.excel_prep.titles import check_titles
from metadata_automation.excel_prep.workbook import (
    class_sheet_names,
    headers,
    property_rows,
    set_value,
    split_iri_suffix,
    text,
)

_CARDINALITY = re.compile(r"^(\d+)(?:\.\.(\d+|n|\*))?$")


def derive_fdp_workbook(src: Path, dst: Path, namespace: str | None = None, overwrite: bool = False) -> Report:
    """Write the FDP variant of ``src`` to ``dst`` and return what was done and what is missing.

    The workbook is first prepared in memory exactly as ``prepare_workbook`` does, plus the ``Resource`` sheet the FDP
    needs; for an already prepared source that changes nothing. Then the FDP rules are applied. ``dst`` must differ
    from ``src`` and must not exist unless ``overwrite``.
    """
    src = Path(src)
    if dst is None:
        raise ValueError("An output path is required; the FDP variant is never written in place.")
    dst = resolve_output(src, Path(dst), overwrite)

    wb = load_workbook(src, rich_text=True)
    report = Report()
    model, untargeted = prepare_in_memory(wb, namespace, report, extra_sheets=FDP_SHEETS)
    if not untargeted:
        apply_fdp_rules(wb, model, report)
        check_titles(wb, model, report)
        add_property_uris(wb, report)

    save_workbook(wb, src, dst, report)
    return report


def apply_fdp_rules(wb: Workbook, model: dict[str, ModelClass], report: Report) -> None:
    """Reference independent classes by IRI in every class sheet."""
    for name in class_sheet_names(wb):
        ws = wb[name]
        cols = headers(ws)
        if "Range" not in cols:
            continue
        for row in property_rows(ws, cols):
            base, has_iri = split_iri_suffix(text(ws, row, cols["Range"]))
            if base in model:
                _apply_model_range(ws, row, cols, base, has_iri, model[base], report)


def _apply_model_range(
    ws: Worksheet,
    row: int,
    cols: dict[str, int],
    base: str,
    has_iri: bool,
    cls: ModelClass,
    report: Report,
) -> None:
    """Independent classes are referenced by IRI with optional cardinality and URI DASH; dependent ones are nested.

    A class without a valid SHACL_independent_class flag is treated as dependent.
    """
    if cls.independent:
        reason = f"'{cls.sheet}' is an independent class"
        if not has_iri:
            set_value(ws, row, cols["Range"], f"{base} (IRI)", report, "Range", reason)
        _relax_cardinality(ws, row, cols, cls, report)
        _reference_by_iri(ws, row, cols, report, reason)
        return
    if has_iri:
        report.issue(ws.title, row, "Range", f"'{cls.sheet}' is a dependent class but the range ends in (IRI)")
    else:
        set_node(ws, row, cols, cls, report, "dependent class")


def _reference_by_iri(ws: Worksheet, row: int, cols: dict[str, int], report: Report, reason: str) -> None:
    """Empty SHACL_sh:node and overwrite the DASH values with the URI viewer and editor."""
    if old := text(ws, row, cols["SHACL_sh:node"]):
        ws.cell(row=row, column=cols["SHACL_sh:node"]).value = None
        report.change(ws.title, row, "SHACL_sh:node", f"cleared '{old}' ({reason})")
    for column, value in zip(("SHACL_dash:viewer", "SHACL_dash:editor"), URI, strict=True):
        if text(ws, row, cols[column]) != value:
            set_value(ws, row, cols[column], value, report, column, reason)


def _relax_cardinality(ws: Worksheet, row: int, cols: dict[str, int], cls: ModelClass, report: Report) -> None:
    if "Cardinality" not in cols:
        return
    current = text(ws, row, cols["Cardinality"])
    relaxed = relax_cardinality(current)
    if relaxed is None:
        report.issue(ws.title, row, "Cardinality", f"cannot parse {current!r}")
    elif relaxed != current:
        reason = f"'{cls.sheet}' is an independent class"
        set_value(ws, row, cols["Cardinality"], relaxed, report, "Cardinality", reason)


def relax_cardinality(cardinality: str) -> str | None:
    """Lower the minimum to 0: ``"1..n"`` -> ``"0..n"``, ``"1"`` -> ``"0..1"``. None when unparsable."""
    match = _CARDINALITY.match(cardinality.strip())
    if not match:
        return None
    low, high = match.group(1), match.group(2)
    if int(low) == 0:
        return cardinality
    return f"0..{high or low}"
