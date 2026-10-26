"""Workbook helpers shared by the preparer and the prefix check."""

import re
from copy import copy

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from metadata_automation.excel_prep.report import Report

NON_CLASS_SHEETS = ("Info", "User Guide", "prefixes", "classes")
IRI_SUFFIX = "(IRI)"

_CURIE = re.compile(r"^([A-Za-z_][\w.-]*):(?!//)(\S*)$")


def class_sheet_names(wb: Workbook) -> list[str]:
    """Every sheet except the general ones is a class sheet."""
    return [name for name in wb.sheetnames if name not in NON_CLASS_SHEETS]


def headers(ws: Worksheet) -> dict[str, int]:
    """Map stripped header names to 1-based column indices."""
    return {str(c.value).strip(): c.column for c in ws[1] if not is_empty(c.value)}


def ensure_columns(ws: Worksheet, columns: list[str], report: Report) -> tuple[dict[str, int], list[str]]:
    """Append missing columns after the last column with data. Returns the headers and the added names.

    Data in columns without a header is reported, as the modeler may have meant it as a pipeline column.
    """
    cols = headers(ws)
    for col in _unheaded_data_columns(ws, cols):
        report.issue(ws.title, None, f"column {get_column_letter(col)}", "data in column without header")
    added = []
    last = _last_data_column(ws)
    for name in columns:
        if name not in cols:
            last += 1
            ws.cell(row=1, column=last, value=name)
            cols[name] = last
            added.append(name)
    return cols, added


def _last_data_column(ws: Worksheet) -> int:
    """Index of the right-most column with a non-empty cell, 0 for an empty sheet."""
    return max((c.column for row in ws.iter_rows() for c in row if not is_empty(c.value)), default=0)


def _unheaded_data_columns(ws: Worksheet, cols: dict[str, int]) -> list[int]:
    """Columns without a header that have a non-empty cell below the header."""
    headed = set(cols.values())
    cells = (c for row in ws.iter_rows(min_row=2) for c in row)
    return sorted({c.column for c in cells if c.column not in headed and not is_empty(c.value)})


def copy_column_format(ws: Worksheet, src_col: int, dst_col: int) -> None:
    """Copy cell styles and width of ``src_col`` onto ``dst_col``, header row included."""
    for row in range(1, ws.max_row + 1):
        src = ws.cell(row=row, column=src_col)
        if src.has_style:
            ws.cell(row=row, column=dst_col)._style = copy(src._style)
    src_dim = ws.column_dimensions[get_column_letter(src_col)]
    if src_dim.width:
        ws.column_dimensions[get_column_letter(dst_col)].width = src_dim.width


def is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def text(ws: Worksheet, row: int, col: int | None) -> str:
    """Cell value as a stripped string, empty when missing."""
    if col is None:
        return ""
    value = ws.cell(row=row, column=col).value
    return "" if is_empty(value) else str(value).strip()


def data_rows(ws: Worksheet, cols: dict[str, int], keys: tuple[str, ...]) -> list[int]:
    """Row numbers below the header where any of the ``keys`` columns has a value."""
    key_cols = [cols[k] for k in keys if k in cols]
    return [r for r in range(2, ws.max_row + 1) if any(text(ws, r, c) for c in key_cols)]


def property_rows(ws: Worksheet, cols: dict[str, int]) -> list[int]:
    return data_rows(ws, cols, ("Property label", "Property URI"))


def split_iri_suffix(value: str) -> tuple[str, bool]:
    """Split ``"dct:Standard (IRI)"`` into ``("dct:Standard", True)``."""
    value = value.strip()
    if value.endswith(IRI_SUFFIX):
        return value[: -len(IRI_SUFFIX)].strip(), True
    return value, False


def curie_prefix(value: str) -> str | None:
    """Prefix of a CURIE, or None for full IRIs and values without a prefix."""
    match = _CURIE.match(value.strip())
    return match.group(1) if match else None


def is_literal(value: str) -> bool:
    """Only rdfs:Literal and the xsd datatypes are literal ranges."""
    value = value.strip()
    return value == "rdfs:Literal" or curie_prefix(value) == "xsd"


def parse_bool(value) -> bool | None:
    """Excel booleans and their text forms; None when empty or unrecognised."""
    if isinstance(value, bool):
        return value
    if is_empty(value):
        return None
    lowered = str(value).strip().lower()
    if lowered in ("true", "1", "yes"):
        return True
    if lowered in ("false", "0", "no"):
        return False
    return None


def property_uri(wb: Workbook, sheet: str, row: int | None) -> str | None:
    """Property URI of a class sheet row, or None for other sheets and sheet-level locations."""
    if not row or sheet not in class_sheet_names(wb):
        return None
    ws = wb[sheet]
    return text(ws, row, headers(ws).get("Property URI")) or None


def set_value(ws: Worksheet, row: int, col: int, value: str, report: Report, column: str, reason: str) -> None:
    """Write ``value`` and report the change with its ``reason``."""
    old = text(ws, row, col)
    ws.cell(row=row, column=col, value=value)
    change = f"'{old}' -> '{value}'" if old else f"set to '{value}'"
    report.change(ws.title, row, column, f"{change} ({reason})")
