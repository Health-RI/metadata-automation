"""Check that prefixes are used consistently and registered in the 'prefixes' sheet."""

from collections import defaultdict
from collections.abc import Iterator

from openpyxl.workbook import Workbook

from metadata_automation.excel_prep.report import Cell, Finding
from metadata_automation.excel_prep.workbook import (
    class_sheet_names,
    curie_prefix,
    data_rows,
    headers,
    property_rows,
    property_uri,
    split_iri_suffix,
    text,
)

CLASS_SHEET_COLUMNS = ("Property URI", "Range", "SHACL_sh:node", "SHACL_dash:viewer", "SHACL_dash:editor")
CLASSES_COLUMNS = ("class_URI", "SHACL_target_ontology_name")
VALUE_CHECK_COLUMNS = ("Property URI", "Range")
MAX_LOCATIONS = 3


def iter_values(wb: Workbook) -> Iterator[tuple[Cell, str]]:
    """Yield the raw (unstripped) text of every cell that may hold a CURIE."""
    for name in class_sheet_names(wb):
        ws = wb[name]
        cols = headers(ws)
        for row in property_rows(ws, cols):
            for column in CLASS_SHEET_COLUMNS:
                yield from _raw(ws, row, cols, column)
    if "classes" in wb.sheetnames:
        ws = wb["classes"]
        cols = headers(ws)
        for row in data_rows(ws, cols, ("sheet_name",)):
            for column in CLASSES_COLUMNS:
                yield from _raw(ws, row, cols, column)


def _raw(ws, row: int, cols: dict[str, int], column: str) -> Iterator[tuple[Cell, str]]:
    if column in cols and text(ws, row, cols[column]):
        yield Cell(ws.title, row, column), str(ws.cell(row=row, column=cols[column]).value)


def used_prefixes(wb: Workbook) -> dict[str, list[Cell]]:
    """Map each prefix used in the workbook to the cells using it."""
    used = defaultdict(list)
    for cell, raw in iter_values(wb):
        prefix = curie_prefix(split_iri_suffix(raw)[0])
        if prefix:
            used[prefix].append(cell)
    return dict(used)


def registered_prefixes(wb: Workbook) -> list[tuple[int, str, str]]:
    """Rows of the 'prefixes' sheet as ``(row, prefix, namespace)``."""
    if "prefixes" not in wb.sheetnames:
        return []
    ws = wb["prefixes"]
    cols = headers(ws)
    return [
        (row, text(ws, row, cols.get("prefix")), text(ws, row, cols.get("namespace")))
        for row in data_rows(ws, cols, ("prefix",))
    ]


def namespace_map(wb: Workbook) -> dict[str, str]:
    """Prefix to namespace, first registration wins."""
    namespaces = {}
    for _row, prefix, namespace in registered_prefixes(wb):
        namespaces.setdefault(prefix, namespace)
    return namespaces


def expand(curie: str, namespaces: dict[str, str]) -> str | None:
    """Full IRI of a CURIE, or None when its prefix has no namespace."""
    prefix = curie_prefix(curie)
    if prefix and namespaces.get(prefix):
        return namespaces[prefix] + curie.split(":", 1)[1]
    return None


def same_term(value: str, curie: str, namespaces: dict[str, str]) -> bool:
    """True when ``value`` is ``curie`` as written, or a CURIE that expands to the same IRI."""
    if value == curie:
        return True
    iri = expand(curie, namespaces)
    return iri is not None and expand(value, namespaces) == iri


def check_prefixes(wb: Workbook) -> list[Finding]:
    """Report unregistered, duplicate, aliased, case-variant and malformed prefix usage."""
    registered = registered_prefixes(wb)
    used = used_prefixes(wb)
    return (
        _duplicates(registered)
        + _aliases(registered)
        + _case_variants(registered, used)
        + _unregistered(wb, registered, used)
        + _malformed_values(wb, registered)
    )


def _duplicates(registered: list[tuple[int, str, str]]) -> list[Finding]:
    by_prefix = defaultdict(list)
    for row, prefix, namespace in registered:
        by_prefix[prefix].append((row, namespace))
    findings = []
    for prefix, entries in by_prefix.items():
        if len(entries) < 2:
            continue
        rows = ", ".join(str(row) for row, _ in entries)
        namespaces = {ns for _, ns in entries}
        detail = f"conflicting namespaces {sorted(namespaces)}" if len(namespaces) > 1 else "the same namespace"
        findings.append(
            Finding(Cell("prefixes", entries[0][0], "prefix"), f"'{prefix}' registered in rows {rows} with {detail}")
        )
    return findings


def _aliases(registered: list[tuple[int, str, str]]) -> list[Finding]:
    by_namespace = defaultdict(list)
    for row, prefix, namespace in registered:
        if namespace and prefix not in [p for _, p in by_namespace[namespace]]:
            by_namespace[namespace].append((row, prefix))
    return [
        Finding(
            Cell("prefixes", entries[0][0], "namespace"),
            f"namespace {namespace} registered under several prefixes: {', '.join(p for _, p in entries)}",
        )
        for namespace, entries in by_namespace.items()
        if len(entries) > 1
    ]


def _case_variants(registered: list[tuple[int, str, str]], used: dict[str, list[Cell]]) -> list[Finding]:
    by_lower = defaultdict(set)
    for prefix in {p for _, p, _ in registered} | set(used):
        by_lower[prefix.lower()].add(prefix)
    return [
        Finding(Cell("prefixes"), f"prefix spelled in different cases: {', '.join(sorted(variants))}")
        for variants in by_lower.values()
        if len(variants) > 1
    ]


def _unregistered(wb: Workbook, registered: list[tuple[int, str, str]], used: dict[str, list[Cell]]) -> list[Finding]:
    namespaces = {}
    for _row, prefix, namespace in registered:
        namespaces.setdefault(prefix, namespace)
    findings = []
    for prefix in sorted(used):
        if namespaces.get(prefix):
            continue
        cells = used[prefix]
        where = "; ".join(_location(wb, c) for c in cells[:MAX_LOCATIONS])
        if len(cells) > MAX_LOCATIONS:
            where += f"; and {len(cells) - MAX_LOCATIONS} more"
        findings.append(Finding(Cell("prefixes"), f"'{prefix}' has no registered namespace, used in: {where}"))
    return findings


def _location(wb: Workbook, cell: Cell) -> str:
    uri = property_uri(wb, cell.sheet, cell.row)
    return f"{cell} ({uri})" if uri else str(cell)


def _malformed_values(wb: Workbook, registered: list[tuple[int, str, str]]) -> list[Finding]:
    namespaces = [(prefix, ns) for _, prefix, ns in registered if ns]
    findings = []
    for cell, raw in iter_values(wb):
        if cell.column not in VALUE_CHECK_COLUMNS:
            continue
        if raw != raw.strip() or "  " in raw:
            findings.append(Finding(cell, f"stray whitespace in {raw!r}"))
        base, _ = split_iri_suffix(raw)
        if base.startswith(("http://", "https://")):
            match = next(((p, ns) for p, ns in namespaces if base.startswith(ns)), None)
            if match:
                prefix, ns = match
                findings.append(Finding(cell, f"full IRI {base!r}, use CURIE '{prefix}:{base[len(ns) :]}'"))
        elif curie_prefix(base) is None:
            findings.append(Finding(cell, f"no prefix in {raw.strip()!r}"))
    return findings
