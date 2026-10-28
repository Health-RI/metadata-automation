"""Add the general sheets, columns and rows, and read the model classes."""

from collections import defaultdict
from dataclasses import dataclass

from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from metadata_automation.excel_prep.columns import CLASSES_COLUMNS, PREFIXES_COLUMNS, SHEET_COLUMNS
from metadata_automation.excel_prep.prefix_check import used_prefixes
from metadata_automation.excel_prep.report import Report
from metadata_automation.excel_prep.workbook import (
    NON_CLASS_SHEETS,
    class_sheet_names,
    curie_prefix,
    data_rows,
    ensure_columns,
    headers,
    parse_bool,
    text,
)


@dataclass
class ModelClass:
    """A class of this metadata model, as configured in the 'classes' sheet."""

    sheet: str
    class_uri: str
    independent: bool | None


def normalize_sheet_names(wb: Workbook, report: Report) -> None:
    """Rename case variants of the general sheets, e.g. 'Classes', to the names the pipeline reads."""
    for name in NON_CLASS_SHEETS:
        for title in wb.sheetnames:
            if title != name and title.lower() == name.lower():
                ws = wb[title]
                # openpyxl counts the sheet's own name as a clash, so go through a free name.
                ws.title = f"{name}__tmp"
                ws.title = name
                report.change(name, None, None, f"renamed from '{title}'")


def add_general_sheets(wb: Workbook, names: list[str], report: Report) -> None:
    """Create the missing ``names`` sheets after Info / User Guide; a case variant counts as present."""
    index = max((wb.sheetnames.index(n) + 1 for n in ("Info", "User Guide") if n in wb.sheetnames), default=0)
    for name in names:
        lowered = [n.lower() for n in wb.sheetnames]
        if name.lower() in lowered:
            index = lowered.index(name.lower()) + 1
            continue
        ws = wb.create_sheet(name, index)
        for col, header in enumerate(SHEET_COLUMNS[name], start=1):
            ws.cell(row=1, column=col, value=header)
        report.change(name, None, None, "added sheet")
        index += 1


def check_duplicates(wb: Workbook) -> None:
    """Raise when the 'classes' sheet repeats a sheet_name or SHACL_target_ontology_name."""
    if "classes" not in wb.sheetnames:
        return
    ws = wb["classes"]
    cols = headers(ws)
    duplicates = []
    for column in ("sheet_name", "SHACL_target_ontology_name"):
        if column not in cols:
            continue
        rows = defaultdict(list)
        for row in data_rows(ws, cols, (column,)):
            rows[text(ws, row, cols[column])].append(row)
        duplicates += [f"{column} '{v}' in rows {', '.join(map(str, r))}" for v, r in rows.items() if len(r) > 1]
    if duplicates:
        raise ValueError("Duplicates in the 'classes' sheet; fix them and run again:\n  " + "\n  ".join(duplicates))


def prepare_classes_sheet(
    wb: Workbook, namespace: str | None, report: Report
) -> tuple[dict[str, ModelClass], list[str]]:
    """Add missing columns and rows, and derive class_URI when ``namespace`` is given.

    Returns the model classes and the class sheets with an empty SHACL_target_ontology_name.
    """
    ws = wb["classes"]
    cols, added = ensure_columns(ws, CLASSES_COLUMNS, report)
    if added:
        report.change("classes", None, None, f"added columns {', '.join(added)}")

    sheets = class_sheet_names(wb)
    rows = _class_rows(ws, cols, sheets, report)

    model: dict[str, ModelClass] = {}
    untargeted = []
    for name in sheets:
        row = rows[name]
        target = text(ws, row, cols["SHACL_target_ontology_name"])
        class_uri = text(ws, row, cols["class_URI"])
        if namespace and not class_uri and curie_prefix(target):
            class_uri = f"{namespace}:{target.split(':', 1)[1]}"
            ws.cell(row=row, column=cols["class_URI"], value=class_uri)
            report.change("classes", row, "class_URI", f"set to '{class_uri}'")

        raw_independent = ws.cell(row=row, column=cols["SHACL_independent_class"]).value
        independent = parse_bool(raw_independent)
        if independent is None and text(ws, row, cols["SHACL_independent_class"]):
            report.issue("classes", row, "SHACL_independent_class", f"expected TRUE or FALSE, got {raw_independent!r}")
        elif independent is None:
            report.issue("classes", row, "SHACL_independent_class", f"empty for '{name}'; treated as FALSE (dependent)")
        elif independent:
            report.issue("classes", row, "SHACL_independent_class", f"TRUE for '{name}'; treated as independent")

        if not target:
            untargeted.append(name)
            continue
        model[target] = ModelClass(name, class_uri, independent)
    return model, untargeted


def _class_rows(ws: Worksheet, cols: dict[str, int], sheets: list[str], report: Report) -> dict[str, int]:
    """Map class sheets to their 'classes' row, adding rows for missing sheets."""
    rows = {text(ws, r, cols["sheet_name"]): r for r in data_rows(ws, cols, ("sheet_name",))}
    for name, row in rows.items():
        if name not in sheets:
            general = name.lower() in {n.lower() for n in NON_CLASS_SHEETS}
            reason = "is not a class sheet" if general else "does not exist"
            report.issue("classes", row, "sheet_name", f"sheet '{name}' {reason}")

    next_row = max(rows.values(), default=1) + 1
    for name in sheets:
        if name not in rows:
            ws.cell(row=next_row, column=cols["sheet_name"], value=name)
            rows[name] = next_row
            report.change("classes", next_row, "sheet_name", f"added row for sheet '{name}'")
            next_row += 1
    return rows


def prepare_prefixes_sheet(wb: Workbook, report: Report) -> None:
    """Add a row, without namespace, for every used prefix that is not registered."""
    ws = wb["prefixes"]
    cols, added = ensure_columns(ws, PREFIXES_COLUMNS, report)
    if added:
        report.change("prefixes", None, None, f"added columns {', '.join(added)}")
    rows = data_rows(ws, cols, ("prefix",))
    registered = {text(ws, r, cols["prefix"]) for r in rows}
    next_row = max(rows, default=1) + 1
    for prefix in sorted(used_prefixes(wb)):
        if prefix not in registered:
            ws.cell(row=next_row, column=cols["prefix"], value=prefix)
            report.change("prefixes", next_row, "prefix", f"added '{prefix}' (namespace unknown)")
            next_row += 1
