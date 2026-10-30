"""Apply a plan to the source Excel.

The plan is checked against the source first (hash, rows, current values); nothing is written when anything does not
fit. Edits marked ``no`` are skipped, ``Value to apply`` on the Gaps sheet is written, and every change is logged.
"""

import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from metadata_automation.excel_prep.columns import SHEET_COLUMNS
from metadata_automation.excel_prep.drawings import drawing_sheets
from metadata_automation.excel_prep.plan import (
    ADD_COLUMN,
    ADD_ROW,
    ADD_SHEET,
    KIND_ORDER,
    RENAME_SHEET,
    SET_CELL,
    Edit,
    Plan,
    read_plan,
    row_key,
    source_hash,
)
from metadata_automation.excel_prep.preparer import resolve_output, save_workbook
from metadata_automation.excel_prep.report import Report
from metadata_automation.excel_prep.structure import add_general_sheets
from metadata_automation.excel_prep.workbook import (
    NON_CLASS_SHEETS,
    copy_column_format,
    data_rows,
    ensure_columns,
    headers,
    text,
)


class PlanError(ValueError):
    """The plan does not fit the source. Carries every problem found."""

    def __init__(self, problems: list[str]) -> None:
        self.problems = problems
        super().__init__("The plan cannot be applied:\n  " + "\n  ".join(problems))


@dataclass
class ApplyResult:
    """What was done, for the log trail."""

    applied: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    filled: list[str] = field(default_factory=list)
    open_gaps: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)

    def format(self) -> str:
        lines = []
        for title, items in (
            ("Applied", self.applied),
            ("Gap values written", self.filled),
            ("Skipped (Apply = no)", self.skipped),
            ("Required cells still empty", self.open_gaps),
            ("Issues", self.issues),
        ):
            lines.append(f"{title} ({len(items)}):")
            lines += [f"  - {item}" for item in items] or ["  none"]
            lines.append("")
        return "\n".join(lines)


def apply_plan(
    src: Path,
    plan_path: Path,
    dst: Path | None = None,
    overwrite: bool = False,
    keep_plan: bool = False,
    check_hash: bool = True,
) -> ApplyResult:
    """Apply the plan at ``plan_path`` to ``src``, in place or to ``dst``.

    With ``check_hash`` False the source may have changed since the plan was made, for instance by rows being added
    or sorted; every edit is then still matched on its key and checked against the value it expected.

    Raises:
        PlanError: When the plan was made from another source, or an edit does not fit. Nothing is written then.
    """
    src, plan_path = Path(src), Path(plan_path)
    plan = read_plan(plan_path)
    if check_hash and plan.source_hash != source_hash(src):
        raise PlanError(
            [f"{src.name} changed since the plan was made; run plan-excel again, or pass --ignore-source-hash."]
        )
    dst = resolve_output(src, dst, overwrite)

    wb = load_workbook(src, rich_text=True)
    result = ApplyResult()
    problems: list[str] = []
    _apply_edits(wb, plan, result, problems)
    _apply_gaps(wb, plan, result, problems)
    if problems:
        raise PlanError(problems)

    report = Report()
    if dst == src and drawing_sheets(src):
        backup = src.with_name(src.name + ".bak")
        shutil.copy2(src, backup)
        result.issues.append(f"copied to {backup.name} before the in-place update, as it has drawings")
    save_workbook(wb, src, dst, report)
    result.issues += [str(f) for f in report.issues]
    if dst == src and not keep_plan:
        plan_path.unlink()
        result.issues.append(f"deleted {plan_path.name}")
    return result


def _apply_edits(wb: Workbook, plan: Plan, result: ApplyResult, problems: list[str]) -> None:
    ordered = sorted(plan.edits, key=lambda e: KIND_ORDER.index(e.kind) if e.kind in KIND_ORDER else len(KIND_ORDER))
    for edit in ordered:
        label = _label(edit)
        if edit.kind not in KIND_ORDER:
            problems.append(f"{label}: unknown kind '{edit.kind}'")
        elif not edit.apply:
            result.skipped.append(label)
        elif error := _apply_edit(wb, edit):
            problems.append(f"{label}: {error}")
        else:
            result.applied.append(label)


def _label(e: Edit) -> str:
    where = f"{e.sheet} / {e.key} / {e.column}" if e.key else f"{e.sheet} / {e.column}" if e.column else e.sheet
    if e.kind == SET_CELL:
        return f"{where}: '{e.old}' -> '{e.new}' ({e.reason})"
    if e.kind == RENAME_SHEET:
        return f"rename sheet '{e.old}' to '{e.new}'"
    return f"{e.kind}: {where}" + (f" ({e.reason})" if e.reason else "")


def _apply_edit(wb: Workbook, edit: Edit) -> str | None:
    """Apply one edit; returns why it does not fit, or None."""
    if edit.kind == RENAME_SHEET:
        return _rename(wb, edit)
    if edit.kind == ADD_SHEET:
        return _add_sheet(wb, edit)
    if edit.sheet not in wb.sheetnames:
        return f"sheet '{edit.sheet}' does not exist"
    ws = wb[edit.sheet]
    if edit.kind == ADD_COLUMN:
        _add_column(ws, edit.column)
        return None
    if edit.kind == ADD_ROW:
        return _add_row(ws, edit)
    return _set_cell(ws, edit.key, edit.occurrence, edit.column, edit.new, expected=edit.old)


def _rename(wb: Workbook, edit: Edit) -> str | None:
    if edit.old not in wb.sheetnames:
        return f"sheet '{edit.old}' does not exist"
    ws = wb[edit.old]
    # openpyxl counts the sheet's own name as a clash, so go through a free name.
    ws.title = f"{edit.new}__tmp"
    ws.title = edit.new
    return None


def _add_sheet(wb: Workbook, edit: Edit) -> str | None:
    if edit.sheet.lower() in [n.lower() for n in wb.sheetnames]:
        return f"sheet '{edit.sheet}' already exists"
    if edit.sheet in SHEET_COLUMNS:
        add_general_sheets(wb, [edit.sheet], Report())
    else:
        wb.create_sheet(edit.sheet)
    return None


def _add_column(ws, column: str) -> None:
    cols, added = ensure_columns(ws, [column], Report())
    if added and ws.title not in NON_CLASS_SHEETS:
        copy_column_format(ws, 1, cols[column])


def _add_row(ws, edit: Edit) -> str | None:
    cols = headers(ws)
    if edit.column not in cols:
        return f"column '{edit.column}' does not exist"
    if locate_row(ws, cols, edit.key, 1) is not None:
        return f"a row with key '{edit.key}' already exists"
    rows = data_rows(ws, cols, (edit.column,))
    ws.cell(row=max(rows, default=1) + 1, column=cols[edit.column], value=edit.new)
    return None


def locate_row(ws, cols: dict[str, int], key: str, occurrence: int) -> int | None:
    """The row of the ``occurrence``-th row with ``key``, None when there is none."""
    matches = [r for r in range(2, ws.max_row + 1) if row_key(ws, r, cols) == key]
    return matches[occurrence - 1] if 0 < occurrence <= len(matches) else None


def _set_cell(ws, key: str, occurrence: int, column: str, value: str, expected: str | None) -> str | None:
    cols = headers(ws)
    if column not in cols:
        return f"column '{column}' does not exist in '{ws.title}'"
    row = locate_row(ws, cols, key, occurrence)
    if row is None:
        return f"no row with key '{key}' (occurrence {occurrence}) in '{ws.title}'"
    current = text(ws, row, cols[column])
    if expected is not None and current != expected:
        return f"current value is '{current}', the plan expected '{expected}'"
    ws.cell(row=row, column=cols[column], value=value or None)
    return None


def _apply_gaps(wb: Workbook, plan: Plan, result: ApplyResult, problems: list[str]) -> None:
    for gap in plan.gaps:
        label = f"{gap.sheet} / {gap.key} / {gap.column}"
        if gap.value in ("", None):
            result.open_gaps.append(label)
            continue
        if gap.sheet not in wb.sheetnames:
            problems.append(f"{label}: sheet '{gap.sheet}' does not exist")
        elif error := _set_cell(wb[gap.sheet], gap.key, gap.occurrence, gap.column, str(gap.value), expected=None):
            problems.append(f"{label}: {error}")
        else:
            result.filled.append(f"{label}: '{gap.value}'")
